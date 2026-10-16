"""Excel-Rechenmodell Baufeld 3 (6 Wohnungen, große Gastro im EG, 960 m² BGF).
Annahmen, Kosten, Kredit (Bauzeit + Aufschub), Cashflow rechnen mit Formeln;
Szenarien sind Werte aus model.js (Stand Oktober 2026)."""
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

szen = json.load(open(sys.argv[1]))
out = sys.argv[2]

GREEN = "2F4A3E"
RED = "8C3B32"
hdr_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
hdr_fill = PatternFill("solid", fgColor=GREEN)
inp_font = Font(name="Calibri", color="0000FF", size=11)  # Eingaben blau
bold = Font(name="Calibri", bold=True, size=11)
base = Font(name="Calibri", size=11)
note = Font(name="Calibri", italic=True, size=10, color="5B6770")
sec_fill = PatternFill("solid", fgColor="F1F0EE")
EUR = '#,##0 "€";[Red]-#,##0 "€"'
EURS = '+#,##0 "€";[Red]-#,##0 "€"'
PCT = "0.0%"

wb = Workbook()


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name="Cambria", bold=True, size=16, color=RED)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = note


def header(ws, row, labels):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=i + 1, value=t)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="left" if i == 0 else "right")


def section(ws, row, text, cols=3):
    for i in range(cols):
        ws.cell(row=row, column=i + 1).fill = sec_fill
    ws.cell(row=row, column=1, value=text).font = bold


# ---------------- Annahmen ----------------
ws = wb.active
ws.title = "Annahmen"
title(ws, "Annahmen", "Blaue Zellen sind Eingaben. Alle anderen Blätter rechnen daraus.")
rows = [
    (3, "Grundstück und Bau", None, None, None),
    (4, "Grundstück (m²)", 960, "#,##0", "Baufeld 3, ca. 960 m²"),
    (5, "Preis je m² (€)", 230, EUR, ""),
    (6, "Grunderwerbsteuer", 0.05, PCT, "Steuer beim Grundstückskauf"),
    (7, "Notar und Grundbuch", 0.02, PCT, "Annahme"),
    (8, "Bruttogrundfläche BGF (m²)", 960, "#,##0", "EG 340 + 1. OG 340 + Staffelgeschoss 280 (Holzbau)"),
    (9, "Baukosten je m² BGF (€)", 2200, EUR, "Annahme des Bauherrn (Bekanntenpreis), Angebot nötig"),
    (10, "Baunebenkosten", 0.15, PCT, "Architekt, Statik, Gebühren"),
    (11, "Gastro-Ausbau (€)", 150000, EUR, "Annahme des Bauherrn, einfacher Ausbau (erste Schätzung 310.000 €)"),
    (12, "Förderung (Zuschuss in % der Kosten)", 0, PCT, "Annahme, mit der Stadt klären (0 bis 25 %)"),
    (14, "Finanzierung", None, None, None),
    (15, "Zins Plan", 0.033, PCT, "Mischzins aus Förderkrediten (Annahme)"),
    (16, "Zins Ausgang", 0.04, PCT, "10 Jahre fest, Stand 17.09.2026"),
    (17, "Zins für diese Rechnung", "=B15", PCT, "Auf =B16 stellen für den Ausgangszins"),
    (18, "Tilgung", 0.01, PCT, "Rückzahlung der Schuld, 0 % = nur Zinsen"),
    (19, "Bauzeit (Monate)", 15, "0", "Plan im Blatt Kredit ist auf 15 Monate ausgelegt"),
    (20, "Bereitstellungszins p. a.", 0.03, PCT, "Gebühr auf zugesagtes, noch nicht abgerufenes Geld"),
    (21, "Bereitstellungszins ab Monat", 7, "0", ""),
    (22, "Raten-Aufschub (Monate)", 12, "0", "Plan im Blatt Kredit ist auf 12 Monate ausgelegt"),
    (24, "Einnahmen und Gastro", None, None, None),
    (25, "Anzahl Wohnungen", 6, "0", "3 im 1. OG, 3 im Staffelgeschoss"),
    (26, "Kaltmiete je Wohnung (€)", 950, EUR, "Annahme des Bauherrn, Markt ca. 11 €/m²"),
    (27, "Gastrofläche (m²)", 290, "#,##0", ""),
    (28, "Pacht je m² (€)", 12, EUR, "Gastro zahlt an das Haus"),
    (29, "Rücklagen", 0.10, PCT, "Reparaturen, Leerstand, in % von Miete und Pacht"),
    (30, "Gastro-Umsatz pro Monat (€)", 50000, EUR, "netto"),
    (31, "Gastro-Marge vor Pacht", 0.305, PCT, "Branchen-Richtwert"),
    (32, "Dein Lohn pro Monat (€)", 0, EUR, "wird vorerst nicht gerechnet"),
    (34, "Holzbau und Terrasse", None, None, None),
    (35, "Staffelgeschoss in Holzbau (m² BGF)", 280, "#,##0", "Zimmerei, Terrasse rundum"),
    (36, "Holzbau günstiger als Massivbau (€ je m²)", 400, EUR, "Annahme, Angebot der Zimmerei nötig"),
    (37, "Dachterrasse rundum (m²)", 58, "#,##0", "zählt nicht zur BGF"),
    (38, "Terrasse Preis (€ je m²)", 450, EUR, "Abdichtung, Holzrost, Geländer, Annahme"),
    (39, "Preisnachlass Bauunternehmer/Zimmerei (%)", 0, PCT, "Bekanntenpreis, noch offen (z. B. 5 % oder 10 %)"),
]
for r, label, val, fmt, hint in rows:
    if val is None:
        section(ws, r, label)
        continue
    ws.cell(row=r, column=1, value=label).font = base
    c = ws.cell(row=r, column=2, value=val)
    c.number_format = fmt
    c.font = inp_font if not (isinstance(val, str) and val.startswith("=")) else base
    ws.cell(row=r, column=3, value=hint).font = note
ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 58

# ---------------- Kosten ----------------
k = wb.create_sheet("Kosten")
title(k, "Kosten bis zum fertigen Haus")
header(k, 3, ["Posten", "Betrag", "Rechnung"])
krows = [
    (4, "Grundstück", "=Annahmen!B4*Annahmen!B5", "m² × Preis"),
    (5, "Grunderwerbsteuer", "=B4*Annahmen!B6", ""),
    (6, "Notar und Grundbuch", "=B4*Annahmen!B7", ""),
    (7, "Grundstück gesamt", "=SUM(B4:B6)", ""),
    (8, "Bau", "=(Annahmen!B8*Annahmen!B9-Annahmen!B35*Annahmen!B36+Annahmen!B37*Annahmen!B38)*(1-Annahmen!B39)", "BGF × Baukosten, minus Holz-Abschlag, plus Terrasse, minus Nachlass"),
    (9, "Baunebenkosten", "=B8*Annahmen!B10", ""),
    (10, "Bau gesamt", "=B8+B9", ""),
    (11, "Gastro-Ausbau", "=Annahmen!B11", ""),
    (12, "Gesamtkosten", "=B7+B10+B11", ""),
    (13, "Förderung (Zuschuss)", "=-B12*Annahmen!B12", ""),
    (14, "Kosten nach Förderung", "=B12+B13", ""),
    (15, "Kosten je m² BGF", "=B14/Annahmen!B8", ""),
    (17, "Grundstück nach Förderung (für den Kredit)", "=B7*(1-Annahmen!B12)", ""),
    (18, "Bau und Gastro nach Förderung (für den Kredit)", "=(B10+B11)*(1-Annahmen!B12)", ""),
]
for r, label, f, h in krows:
    k.cell(row=r, column=1, value=label).font = bold if label in ("Gesamtkosten", "Kosten nach Förderung") else base
    c = k.cell(row=r, column=2, value=f)
    c.number_format = EUR
    c.font = bold if label in ("Gesamtkosten", "Kosten nach Förderung") else base
    k.cell(row=r, column=3, value=h).font = note
k.column_dimensions["A"].width = 46
k.column_dimensions["B"].width = 18
k.column_dimensions["C"].width = 30

# ---------------- Kredit ----------------
c_ = wb.create_sheet("Kredit")
title(c_, "Kredit: Bauzeit, Aufschub und Rate", "Zinsen und Bereitstellungszinsen werden dem Kredit zugeschlagen (kein Eigenkapital).")
header(c_, 4, ["Monat", "Auszahlung", "Zinsen", "Bereitstellung", "Noch nicht abgerufen", "Kredit am Monatsende"])
c_["A5"] = 0
c_["B5"] = "=Kosten!B17"
c_["C5"] = 0
c_["D5"] = 0
c_["E5"] = "=Kosten!B18"
c_["F5"] = "=B5"
for m in range(1, 16):
    r = 5 + m
    c_.cell(row=r, column=1, value=m)
    c_.cell(row=r, column=2, value="=Kosten!$B$18/Annahmen!$B$19")
    c_.cell(row=r, column=3, value=f"=F{r-1}*Annahmen!$B$17/12")
    c_.cell(row=r, column=4, value=f"=IF(A{r}>=Annahmen!$B$21,E{r}*Annahmen!$B$20/12,0)")
    c_.cell(row=r, column=5, value="=Kosten!B18" if m == 1 else f"=E{r-1}-B{r-1}")
    c_.cell(row=r, column=6, value=f"=F{r-1}+C{r}+D{r}+B{r}")
c_["A22"] = "Raten-Aufschub: nur Zinsen laufen auf"
c_["A22"].font = bold
header(c_, 23, ["Monat", "", "Zinsen", "", "", "Kredit am Monatsende"])
for m in range(1, 13):
    r = 23 + m
    c_.cell(row=r, column=1, value=m)
    c_.cell(row=r, column=3, value=f"=F{r-1 if m > 1 else 20}*Annahmen!$B$17/12")
    c_.cell(row=r, column=6, value=f"=F{r-1 if m > 1 else 20}+C{r}")
# Ergebnisse
c_["H4"] = "Ergebnis"
c_["H4"].font = bold
res = [
    (5, "Zinsen in der Bauzeit", "=SUM(C6:C20)"),
    (6, "Bereitstellungszinsen", "=SUM(D6:D20)"),
    (7, "Zinsen im Aufschub-Jahr", "=SUM(C24:C35)"),
    (8, "Kredit bei Start der Rate", "=F35"),
    (9, "Rate pro Monat", "=F35*(Annahmen!B17+Annahmen!B18)/12"),
    (10, "   davon Zinsen", "=F35*Annahmen!B17/12"),
    (11, "   davon Tilgung", "=F35*Annahmen!B18/12"),
]
for r, label, f in res:
    c_.cell(row=r, column=8, value=label).font = bold if r in (8, 9) else base
    cell = c_.cell(row=r, column=9, value=f)
    cell.number_format = EUR
    cell.font = bold if r in (8, 9) else base
for r in range(5, 36):
    for col in (2, 3, 4, 5, 6):
        c_.cell(row=r, column=col).number_format = EUR
c_.column_dimensions["A"].width = 10
for col in "BCDEF":
    c_.column_dimensions[col].width = 20
c_.column_dimensions["G"].width = 4
c_.column_dimensions["H"].width = 30
c_.column_dimensions["I"].width = 18

# ---------------- Cashflow ----------------
cf = wb.create_sheet("Cashflow")
title(cf, "Was pro Monat übrig bleibt", "Mit den Annahmen aus dem Blatt Annahmen (Zins, Förderung, Umsatz).")
header(cf, 4, ["Posten", "pro Monat"])
cfrows = [
    (5, "Wohnungen (Kaltmiete)", "=Annahmen!B25*Annahmen!B26"),
    (6, "Gastro-Pacht", "=Annahmen!B27*Annahmen!B28"),
    (7, "Miete und Pacht gesamt", "=B5+B6"),
    (8, "Gastro-Gewinn nach Pacht", "=Annahmen!B30*Annahmen!B31-B6"),
    (9, "Rate (Zinsen und Tilgung)", "=-Kredit!I9"),
    (10, "Dein Lohn", "=-Annahmen!B32"),
    (11, "Rücklagen", "=-B7*Annahmen!B29"),
    (12, "Ergebnis pro Monat", "=B5+B6+B8+B9+B10+B11"),
    (13, "Ergebnis pro Jahr", "=B12*12"),
    (15, "Break-even Gastro-Umsatz (Ergebnis = 0, abzüglich Lohn aus Annahmen)", "=(Annahmen!B32+Kredit!I9+B7*Annahmen!B29-B7+B6)/Annahmen!B31"),
]
for r, label, f in cfrows:
    cf.cell(row=r, column=1, value=label).font = bold if r in (12, 15) else base
    c = cf.cell(row=r, column=2, value=f)
    c.number_format = EURS if r in (12, 13) else EUR
    c.font = bold if r in (12, 15) else base
cf["A17"] = "Die Tilgung in der Rate ist Vermögensaufbau (die Schuld sinkt), kein Verlust."
cf["A17"].font = note
cf.column_dimensions["A"].width = 52
cf.column_dimensions["B"].width = 18

# ---------------- Szenarien (Werte) ----------------
s = wb.create_sheet("Szenarien")
title(s, "Szenarien (Werte aus dem Rechenmodell, Stand Oktober 2026)", "Gastro-Umsatz 50.000 €, Mieten 6 × 950 €. Ändern sich die Annahmen, bitte neu erzeugen.")
header(s, 4, ["Förderung", "Gesamtkosten", "Rate Plan 3,3 %", "Ergebnis Plan", "Rate Ausgang 4,0 %", "Ergebnis Ausgang", "Break-even Umsatz Plan", "Break-even Umsatz Ausgang"])
for i, d in enumerate(szen["foerderung"]):
    r = 5 + i
    vals = [d["f"], d["gesamt"], d["ratePlan"], d["cfPlan"], d["rateAus"], d["cfAus"], d["bePlan"], d["beAus"]]
    fmts = [PCT, EUR, EUR, EURS, EUR, EURS, EUR, EUR]
    for j, (v, fm) in enumerate(zip(vals, fmts)):
        c = s.cell(row=r, column=j + 1, value=round(v, 2) if j else v)
        c.number_format = fm
s["A12"] = "Ergebnis pro Monat (ohne Lohn des Bauherrn), Plan 3,3 %, je Gastro-Umsatz"
s["A12"].font = bold
header(s, 13, ["Förderung", "40.000 €", "45.000 €", "50.000 €", "55.000 €", "60.000 €"])
for i, d in enumerate(szen["umsatz"]):
    r = 14 + i
    c = s.cell(row=r, column=1, value=d["f"])
    c.number_format = PCT
    for j, v in enumerate(d["werte"]):
        cc = s.cell(row=r, column=2 + j, value=round(v, 2))
        cc.number_format = EURS
for col in "ABCDEFGH":
    s.column_dimensions[col].width = 22

# ---------------- Verkauf je Partei ----------------
vk = wb.create_sheet("Verkauf je Partei")
title(vk, "Verkauf an die Familie: Rechnung je Partei", "Jede Partei mit eigenem Kredit ohne Eigenkapital. Preis = Selbstkosten, 3,3 % Zins, 1 % Tilgung, 10 % Rücklagen. Feste Werte aus model.js.")
header(vk, 4, ["", "A: 1. OG 85 m²", "A: Staffel 63 m²", "A: Gastro (Du)", "B: 1. OG 85 m²", "B: Staffel 63 m²", "B: Gastro + Büro (Du)", "C: Regelgeschoss 85 m²", "C: Staffel 63 m²", "C: Gastro + Büro (Du)"])
VA, VB, VC = szen["verkauf"]["A"], szen["verkauf"]["B"], szen["verkauf"]["C"]
cols = [(VA["og85"], 2), (VA["st56"], 3), (VA["gastro"], 4), (VB["og85"], 5), (VB["st56"], 6), (VB["gastro"], 7), (VC["og85"], 8), (VC["st56"], 9), (VC["gastro"], 10)]
labels = ["Preis (Selbstkosten)", "Rate pro Monat", "Einnahme nach Rücklagen / Gastro-Gewinn (+ Büro)", "Ergebnis pro Monat", "Preis für Ergebnis 0 €", "Lücke zum Selbstkostenpreis"]
for i, lab in enumerate(labels):
    vk.cell(row=5 + i, column=1, value=lab).font = base
for d, c in cols:
    if "kosten" in d and "m2" in d:
        vals = [d["kosten"], d["rate"], d["miete"] * 0.9, d["ergebnis"], d["preisFuerPlus"], d["luecke"]]
    else:
        vals = [d["kosten"], d["rate"], d["gewinn"] + d["buero"], d["ergebnis"], None, None]
    for i, x in enumerate(vals):
        if x is None:
            continue
        cc = vk.cell(row=5 + i, column=c, value=round(x, 2))
        cc.number_format = EURS if i == 3 else EUR
vk.column_dimensions["A"].width = 52
for col in "BCDEFGHIJ":
    vk.column_dimensions[col].width = 20
vb_ = szen["varianteB"]
vk.cell(row=13, column=1, value="Haus als Ganzes mit einem Kredit").font = bold
header(vk, 14, ["", "A: 340 m²", "B: 540 m²", "C: 540 m², 3 Vollgeschosse + Holz-Penthouse"])
for i, (lab, ka) in enumerate([("Gesamtkosten", "gesamt"), ("Rate pro Monat", "rate"), ("Ergebnis pro Monat bei 50.000 € Umsatz", "cf"), ("Break-even Gastro-Umsatz", "be")]):
    vk.cell(row=15 + i, column=1, value=lab).font = base
    for j, key in enumerate(["A", "B", "C"]):
        cc = vk.cell(row=15 + i, column=2 + j, value=round(vb_[key][ka], 2))
        cc.number_format = EURS if ka == "cf" else EUR
vk.cell(row=20, column=1, value="B: EG 540 m² (Gastro 290 + Büro 150), 1. OG 540 m², Staffel 420 m² mit 120 m² Terrasse. C: zusätzlich 2. OG 540 m², Höhe ca. 13 m. Nur mit Ausnahme der Stadt (Baufenster, II Vollgeschosse, 12 m).").font = note

# Ziel: 250 € Plus je Wohnung
zr = szen["ziel"]
vk.cell(row=22, column=1, value="Ziel: 250 € Plus je Wohnung (Käufer ohne Eigenkapital, ohne Tilgung, Variante A)").font = bold
header(vk, 23, ["", "1. OG 85 m²", "Staffel 63 m²", "Du", "Zins 4,0 %: 1. OG", "Zins 4,0 %: Staffel", "Zins 4,0 %: Du"])
zl = ["Verkaufspreis", "Selbstkosten", "Nachlass (von dir getragen)", "Rate pro Monat", "Ergebnis pro Monat"]
for i, lab in enumerate(zl):
    vk.cell(row=24 + i, column=1, value=lab).font = base
for off, key in ((0, "plan"), (3, "ausgang")):
    d = zr[key]
    for j, e in enumerate((d["og"], d["st"])):
        for i, x in enumerate((e["preis"], e["kosten"], e["nachlass"], e["rate"], e["ergebnis"])):
            c_ = vk.cell(row=24 + i, column=2 + j if off == 0 else 5 + j, value=round(x, 2))
            c_.number_format = EURS if i == 4 else EUR
    cc_ = vk.cell(row=26, column=4 if off == 0 else 7, value=round(d["luecke"], 2)); cc_.number_format = EUR
    cc_ = vk.cell(row=27, column=4 if off == 0 else 7, value=round(d["rateEG"] + d["zinsLuecke"], 2)); cc_.number_format = EUR
    cc_ = vk.cell(row=28, column=4 if off == 0 else 7, value=round(d["bessam"], 2)); cc_.number_format = EURS

# Verkaufsplan
vk.cell(row=31, column=1, value="Verkaufsplan (Vorschlag), ohne Förderung, 2.200 €/m², tilgungsfrei, 3,3 % Zins, Notar 2 %, Grunderwerbsteuer 5 % (Vater frei)").font = bold
header(vk, 32, ["Partei", "Wohnung", "Verkaufspreis", "Selbstkosten", "Differenz (von dir getragen)", "Notar, Steuer", "Rate (nur Zinsen)", "Ergebnis pro Monat", "Ergebnis zum Selbstkostenpreis"])
for i, (r, r0) in enumerate(zip(szen["verkaufsplan"], szen["verkaufsplan0"])):
    row = 33 + i
    vk.cell(row=row, column=1, value=r["name"]).font = base
    vk.cell(row=row, column=2, value="1. OG, 85 m²" if r["typ"] == "OG" else "Staffel, 63 m²").font = base
    for col, key in ((3, "preis"), (4, "kosten"), (5, "nachlass"), (6, "nk"), (7, "rate")):
        c_ = vk.cell(row=row, column=col, value=round(r[key], 2)); c_.number_format = EUR
    c_ = vk.cell(row=row, column=8, value=round(r["ergebnis"], 2)); c_.number_format = EURS
    c_ = vk.cell(row=row, column=9, value=round(r0["ergebnis"], 2)); c_.number_format = EURS
sm = szen["verkaufsplanSumme"]
vk.cell(row=40, column=1, value="Differenz gesamt (von dir getragen)").font = base
c_ = vk.cell(row=40, column=3, value=round(sm["luecke"], 2)); c_.number_format = EUR
vk.cell(row=41, column=1, value="Zinsen auf die Differenz pro Monat").font = base
c_ = vk.cell(row=41, column=3, value=round(sm["zinsLuecke"], 2)); c_.number_format = EUR
vk.cell(row=42, column=1, value="Dein Ergebnis pro Monat (Gastro − Rate EG − Zinsen Differenz + eigene Wohnung)").font = bold
c_ = vk.cell(row=42, column=3, value=round(sm["bessam"], 2)); c_.number_format = EURS

# ---------------- Variante 4 Ebenen ----------------
v4 = wb.create_sheet("Variante 4 Ebenen")
title(v4, "Variante: Mittelgeschoss doppelt (4 Ebenen)", "Nur als Gedankenspiel: Höhe ca. 13 m, der B-Plan erlaubt 12 m.")
header(v4, 4, ["", "3 Ebenen", "4 Ebenen"])
ref = szen["referenz"]
v = szen["variante4"]
data = [
    ("Wohnungen", 6, 9, "0"),
    ("BGF (m²)", 960, 1300, "#,##0"),
    ("Gesamtkosten", ref["gesamt"], v["gesamt"], EUR),
    ("Rate (Plan)", ref["rate"], v["rate"], EUR),
    ("Ergebnis pro Monat ohne Förderung", ref["cf"], v["cf0"], EURS),
    ("Ergebnis pro Monat mit 15 % Förderung", szen["foerderung"][3]["cfPlan"], v["cf15"], EURS),
]
for i, (a, b, c, fm) in enumerate(data):
    r = 5 + i
    v4.cell(row=r, column=1, value=a)
    for j, x in enumerate((b, c)):
        cc = v4.cell(row=r, column=2 + j, value=round(x, 2))
        cc.number_format = fm
v4.column_dimensions["A"].width = 44
v4.column_dimensions["B"].width = 18
v4.column_dimensions["C"].width = 18

# ---------------- Hinweise ----------------
h = wb.create_sheet("Hinweise")
title(h, "Hinweise und offene Punkte")
lines = [
    "Alle Zahlen sind Planungsannahmen, keine Zusagen von Bank oder Stadt.",
    "Nebenkosten (Grundsteuer, Gebäudeversicherung, Betriebskosten) werden auf die Mieter umgelegt. Hausverwaltung und Instandhaltung sind nicht umlagefähig und stecken in der Rücklage von 10 %.",
    "Miete 950 € je Wohnung liegt über den Portalwerten für Neubau in Werlte (ca. 10,70–11,02 €/m²).",
    "Förderung ist als Zuschuss in Prozent der Gesamtkosten angesetzt, Höhe und Zweck sind offen.",
    "Gastro-Marge (30,5 % vor Pacht) und Kostenstruktur sind Branchen-Richtwerte, nicht aus Werlte.",
    "Staffelgeschoss (ca. 280 m², Holzbau, schmale Terrasse rundum) und Gesamthöhe ca. 9,90 m müssen mit der Stadt abgestimmt werden.",
    "Zinsbindung 10 Jahre, der Folgezins ist offen.",
]
for i, t in enumerate(lines):
    h.cell(row=3 + i, column=1, value="• " + t).font = base
h.column_dimensions["A"].width = 110

wb.save(out)
print("gespeichert", out)
