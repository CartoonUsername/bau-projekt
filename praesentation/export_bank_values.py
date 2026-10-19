"""Liest die berechneten Werte aus Bankunterlagen_Baufeld3.xlsx (nach recalc.py) und schreibt bank_values.json für Präsentation und Word."""
import json
import sys
import openpyxl

src = sys.argv[1] if len(sys.argv) > 1 else "Bankunterlagen_Baufeld3.xlsx"
reg = json.load(open("bank_registry.json", encoding="utf-8"))
wb = openpyxl.load_workbook(src, data_only=True)
out = {}
for k, ref in reg.items():
    sh, c = ref.split("!")
    out[k] = wb[sh][c].value
# Tabellen
ws = wb["Flächen und Mieten"]
out["einheiten"] = [{"nr": ws[f"A{r}"].value, "geschoss": ws[f"B{r}"].value, "m2": ws[f"C{r}"].value, "nutzung": ws[f"D{r}"].value, "miete_m2": ws[f"E{r}"].value, "miete": ws[f"F{r}"].value, "jahr": ws[f"G{r}"].value} for r in range(12, 20)]
ws = wb["Tilgung Jahre"]
out["tilgung"] = [{"jahr": ws[f"A{r}"].value, "zins": ws[f"B{r}"].value, "tilg": ws[f"C{r}"].value, "rate": ws[f"D{r}"].value, "rest": ws[f"E{r}"].value, "kfw": ws[f"F{r}"].value} for r in range(5, 40)]
ws = wb["Prognose"]
out["prognose"] = {ws[f"A{r}"].value: [ws.cell(row=r, column=2 + y).value for y in range(10)] for r in range(5, 14)}
ws = wb["Ertragswert"]
out["ertrag_tab"] = [[ws[f"A{r}"].value, ws[f"B{r}"].value] for r in list(range(5, 22)) + list(range(25, 37))]
ws = wb["Sensitivität"]
out["sens"] = [[ws[f"A{r}"].value, ws[f"B{r}"].value, ws[f"C{r}"].value, ws[f"D{r}"].value] for r in range(5, 12)]
ws = wb["Kosten und Finanzierung"]
out["kosten_tab"] = [[ws[f"A{r}"].value, ws[f"B{r}"].value] for r in range(5, 19)]
out["fin_tab"] = [[ws[f"A{r}"].value, ws[f"B{r}"].value, ws[f"C{r}"].value, ws[f"D{r}"].value] for r in range(24, 30)]
ws = wb["Gastro Betrieb"]
out["gastro_tab"] = [[ws[f"A{r}"].value, ws[f"B{r}"].value, ws[f"C{r}"].value] for r in range(5, 21)]
out["gastro_umsatz"] = [[ws[f"A{r}"].value, ws[f"C{r}"].value, ws[f"D{r}"].value] for r in range(24, 29)]
ws = wb["Airbnb Prüfung"]
out["airbnb_tab"] = [[ws[f"A{r}"].value] + [ws[f"{c}{r}"].value for c in "BCD"] for r in range(10, 21)]
ws = wb["Bauphase"]
out["bau"] = [[ws.cell(row=r, column=c).value for c in (1, 3, 11, 12, 13, 14)] for r in range(6, 34)]
ws = wb["Bankunterlagen"]
out["unterlagen"] = [[ws[f"A{r}"].value, ws[f"B{r}"].value, ws[f"C{r}"].value, ws[f"D{r}"].value] for r in range(5, 29)]
json.dump(out, open("bank_values.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)
print("bank_values.json", len(out), "Einträge")
