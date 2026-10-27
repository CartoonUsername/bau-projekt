"""Kosten-Nutzen-pro-Kopf (PDF, 1 Seite) und Businessplan (DOCX) je Variante A/B – gleiche Zahlen wie die Präsentation (finanz.py, Planungsannahmen)."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_deck as BD
from build_deck import optdata, TXT, eur, num, pct, mio, F, HERE, hf
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, ListFlowable, ListItem
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import mm
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
GREEN = colors.HexColor("#2D4A3E"); BRICK = colors.HexColor("#8C3B2F")
def pk_pdf(group):
    T = TXT[group]; O = optdata(group); names = T["names"]; multi = len(O) > 1
    ss = getSampleStyleSheet(); h1 = ParagraphStyle("h1", parent=ss["Title"], fontName="Times-Roman", fontSize=20, textColor=BRICK, alignment=0, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, textColor=GREEN, spaceBefore=7, spaceAfter=3); b = ParagraphStyle("b", parent=ss["BodyText"], fontName="Helvetica", fontSize=8.6, leading=11)
    sm = ParagraphStyle("sm", parent=b, fontSize=7.4, textColor=colors.grey)
    def tbl(rows, widths, boldrows=()):
        t = Table(rows, colWidths=widths); st = [("BACKGROUND", (0, 0), (-1, 0), GREEN), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8), ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.lightgrey), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
        for r_ in boldrows: st.append(("FONTNAME", (0, r_), (-1, r_), "Helvetica-Bold"))
        t.setStyle(TableStyle(st)); return t
    out = os.path.join(HERE, f"Kosten_Nutzen_pro_Kopf_{group}.pdf"); doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=12 * mm, bottomMargin=10 * mm)
    S = [Paragraph("NEUE MITTE WERLTE · BAUFELD 3", ParagraphStyle("k", parent=b, textColor=colors.grey, fontSize=8)), Paragraph("Kosten, Nutzen und Gewinn pro Kopf – " + ("Standardversion (B3)" if group == "A" else "überdachte Stellplätze, größere Wohnungen"), h1),
         Paragraph(f"Kurze Erklärung für die Gesellschafter einer GbR. Basis: Rechenmodell mit den Annahmen des früheren Bankplans ({'16 Wohnungen, Büro, Gastro 440 m²' if group == 'A' else 'Optionen B1–B3: 13–14 Wohnungen, Büro, Gastro 376 m²'}), Jahr 1 nach der Anlaufzeit. Alle Zahlen sind Planungsannahmen, Stand 27.10.2026; keine Aussage zur Zulässigkeit.", b)]
    S.append(Paragraph("1 Was kostet das Projekt?", h2))
    S.append(Paragraph("Gesamtkosten: " + " / ".join(eur(r["gesamt"]) for _, _, r, _ in O) + (f" ({', '.join(names)})" if multi else "") + ". Sie werden zu 100 % über KfW-Kredite und Hausbank finanziert, die Gesellschafter zahlen kein Eigenkapital ein. Rate für Zins und 1 % Tilgung: " + " / ".join(f"{eur(r['rate_m'])} im Monat ({eur(r['rate_j'])} im Jahr)" for _, _, r, _ in O) + ".", b))
    S.append(Paragraph("2 Was kommt rein? (pro Jahr)", h2))
    rows = [["Position"] + names] + [["Mieten Wohnen + Büro + Gastro-Pacht"] + [eur(r["roh"]) for _, _, r, _ in O], ["abzüglich Rücklagen (10 %)"] + [eur(-r["rueck"]) for _, _, r, _ in O], ["Gastro-Betriebsergebnis nach Pacht (ohne Lohn)"] + [eur(r["gastro_nach"]) for _, _, r, _ in O], ["Ertrag vor Kapitaldienst"] + [eur(r["vor"]) for _, _, r, _ in O], ["Zins und Tilgung an die Banken"] + [eur(-r["rate_j"]) for _, _, r, _ in O], ["Cash für die Gesellschafter"] + [eur(r["cash"]) for _, _, r, _ in O]]
    S.append(tbl(rows, [80 * mm] + [(178 - 80) * mm / len(O)] * len(O) if False else [70 * mm] + [38 * mm] * len(O), (4, 6)))
    S.append(Spacer(1, 3)); S.append(Paragraph(f"Nutzen-Faktor: Der Ertrag vor Kapitaldienst entspricht " + " / ".join(pct(r["rendite_ges"], 1) for _, _, r, _ in O) + " der Gesamtkosten (Anfangsrendite); nur aus Mieten und Pacht wären es " + " / ".join(pct(r["rendite_imm"], 1) for _, _, r, _ in O) + ", das reicht nicht für die Raten. Der Gastro-Umsatz ist proportional zu den Innenplätzen angesetzt (" + " / ".join(eur(r["umsatz"]) for _, _, r, _ in O[:1]) + " im Monat, ohne Marktbeleg).", b))
    S.append(Paragraph("3 Was bekommt jeder? (gleiche Anteile)", h2))
    if multi:
        rows = [["Personen"] + sum([[f"{n_.split()[0]} Cash/Mon.", f"{n_.split()[0]} nach Steuer"] for n_ in names], [])]
        for p in (1, 2, 4, 8): rows.append([str(p)] + sum([[eur(r["cash"] / 12 / p), eur(r["nach_steuer"] / 12 / p)] for _, _, r, _ in O], []))
        S.append(tbl(rows, [20 * mm] + [26 * mm] * 6)); S.append(Spacer(1, 3))
        S.append(Paragraph("Zu versteuern pro Jahr (gesamt): " + " / ".join(eur(r["gewinn_steuer"]) for _, _, r, _ in O) + "; Tilgung pro Jahr (gesamt, Vermögensaufbau): " + " / ".join(eur(r["tilg1"]) for _, _, r, _ in O) + ".", b))
    else:
        for (k, v, r, d), n_ in zip(O, names):
            rows = [["Personen", "Cash / Monat", "Cash / Jahr", "zu versteuern / Jahr", "nach Steuer / Monat", "Tilgung / Jahr"]]
            for p in (1, 2, 4, 8): rows.append([str(p), eur(r["cash"] / 12 / p), eur(r["cash"] / p), eur(r["gewinn_steuer"] / p), eur(r["nach_steuer"] / 12 / p), eur(r["tilg1"] / p)])
            S.append(tbl(rows, [42 * mm, 28 * mm, 28 * mm, 32 * mm, 30 * mm, 24 * mm])); S.append(Spacer(1, 3))
    r = O[0][2]
    S.append(Paragraph("Steuerlich ist der Gewinn deutlich kleiner als der Cash: Zinsen und Gebäudeabschreibung 3 % mindern ihn, die Tilgung nicht (sie ist Vermögensaufbau). Steuersatz 35 % ist eine Annahme und hängt vom persönlichen Einkommen ab; negative Werte bedeuten Verlust.", b))
    S.append(Paragraph("4 Wo liegt das Risiko?", h2))
    ohne = [((r["wohn"] + r["buero"]) * 0.9 - r["rate_j"]) for _, _, r, _ in O]
    items = [f"Ohne Gastro-Gewinn fehlen " + " / ".join(eur(-x) for x in ohne) + " im Jahr (Kapitaldienstdeckung nur " + " / ".join(num((r['wohn'] + r['buero']) * 0.9 / r['rate_j'], 2) for _, _, r, _ in O) + "): der Gewinn hängt an einer gut laufenden Gastro.", "Der Beleihungsauslauf liegt bei " + " / ".join(pct(r["auslauf"]) for _, _, r, _ in O) + ": die Bank verlangt zusätzliche Sicherheiten (Grundschulden, Bürgschaften oder Eigenkapital); die Lücke zur 80-%-Grenze beträgt " + " / ".join(mio(r["ueber80"]) for _, _, r, _ in O) + ".",
             "Eine GbR haftet mit dem persönlichen Vermögen aller Gesellschafter. Betreibt die GbR auch die Gastro, ist sie komplett gewerblich (Gewerbesteuer); besser ist eine getrennte Betreiberfirma, die Pacht zahlt.", "Lohn des Bauherrn, mögliche Sonderabschreibungen und Gewerbesteuer sind nicht gerechnet. Steuerberater und Bank müssen das prüfen." + (" Bei B kommen offene Punkte zu Grundstück/Baulinie (Vorbau 1,5 m über die Hypothese-Grenze), Kragarmen und Höhenbezug hinzu." if group == "B" else "")]
    S.append(ListFlowable([ListItem(Paragraph(i, b), leftIndent=8) for i in items], bulletType="bullet", start="•", leftIndent=8))
    S.append(Spacer(1, 3)); S.append(Paragraph("Quelle: businessplan/finanz.py (Annahmen wie früherer Bankplan, Nachrechnung dort 0,2 % Abweichung bei den Gesamtkosten).", sm))
    doc.build(S); return out
def plan_docx(group):
    T = TXT[group]; O = optdata(group); names = T["names"]; multi = len(O) > 1; P = F.P
    d = docx.Document(); st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    for sec in d.sections: sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2)
    def H(t):
        h = d.add_heading(t, 1)
        for r in h.runs: r.font.color.rgb = RGBColor(0x8C, 0x3B, 0x2F)
    def PR(t, bold=False):
        p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; return p
    def TB(rows, widths=None, boldrows=()):
        t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j); c.text = ""; r = c.paragraphs[0].add_run(str(val)); r.font.size = Pt(8.5); r.bold = (i == 0 or i in boldrows)
                if j > 0 and i > 0 and str(val)[:1] in "+−-0123456789": c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        d.add_paragraph()
    p = d.add_paragraph(); r = p.add_run("NEUE MITTE WERLTE"); r.bold = True; r.font.size = Pt(9)
    d.add_heading("Businessplan: " + T["titel"].replace("Baufeld 3: ", ""), 0); PR("Baufeld 3, Neue Mitte Werlte · Bessam Hasso, Werlte · Stand 27.10.2026")
    d.add_picture(BD.neu(O[1][0], "hero") if multi else hf("A", "persp"), width=Cm(15.5)); PR("KI-Visualisierung (Nano Banana Pro auf Basis des Geometriemodells) zur Orientierung, keine Planung." + (f" Dargestellt: {names[min(1, len(O)-1)]}." if multi else ""))
    PR("Alle Zahlen sind Planungsannahmen (finanz.py, Annahmen des früheren Bankplans). Mieten, Pacht, Umsatz, Baukosten und KfW-Konditionen müssen durch Angebote, Marktdaten und die Bank bestätigt werden. Keine Aussage zur baurechtlichen Zulässigkeit; keine Variante ist festgelegt.")
    H("1  Zusammenfassung"); PR(T["kurz"] + ".")
    rows = [["Kennzahl"] + names, ["Gesamtkosten"] + [eur(r["gesamt"]) for _, _, r, _ in O], ["Darlehen KfW / Hausbank"] + [f"{eur(r['kfw_w'] + r['kfw_g'])} / {eur(r['bank'])}" for _, _, r, _ in O], ["Kapitaldienst pro Monat"] + [eur(r["rate_m"]) for _, _, r, _ in O], ["Ergebnis pro Monat nach Kapitaldienst"] + [eur(r["cash"] / 12, True) for _, _, r, _ in O],
            ["DSCR gesamt / nur Immobilie"] + [f"{num(r['dscr_ges'],2)} / {num(r['dscr_imm'],2)}" for _, _, r, _ in O], ["Ertragswert / Beleihungsauslauf"] + [f"{mio(r['ew'])} / {pct(r['auslauf'])}" for _, _, r, _ in O]]
    TB(rows)
    H("2  Vorhaben und Standort"); PR("Baufeld 3 östlich des Rathauses, südlich der Planstraße; Fläche ca. 960 m² (Projektangabe, Zuschnitt nicht vermessen). Planungsrecht: B-Plan Nr. 123 (14.04.2023, rechtskräftig, vermutlich MI 3 / Flurstück 59/1): II bis III Vollgeschosse, GRZ 0,6, Traufhöhe 9,80 m, Höhe 15,00 m, Hauptdach ≥ 38°, Bezug OK Fahrbahn. 1. Änderung (Entwurf 25.08.2026): II zwingend, TH 6,50 m, H 12,00 m, erfasst MI 3 nach vorliegendem Stand nicht. Zu klären: Zuordnung Baufeld 3, Baufenster (≈ 27,7 × 15,9 m abgeleitet), Dachform, Höhe, Höhenbezug.")
    H("3  Baukonzept und Flächen")
    rows = [["Fläche"] + names] + [[lab] + [f"{num(d_['flaeche'][kk])} m²" for _, _, _, d_ in O] for lab, kk in (("Erdgeschoss", "eg"), ("1. Obergeschoss", "og1"), ("oberstes Geschoss (Holzbau)", "st"))] + [["BGF gesamt"] + [f"{num(v['bgf'])} m²" for _, v, _, _ in O], ["Wohnfläche netto / Wohnungen"] + [f"{num(v['wfl'])} m² / {d_['mix']['standard']['we']}" for _, v, _, d_ in O]]
    TB(rows, boldrows=(4,)); d.add_picture(os.path.join(BD.GEO, O[min(1, len(O) - 1)][0], "plaene", "schnitt.png"), width=Cm(15.5)); PR("Schnitt schematisch. Geschosshöhen " + ("aus dem B3-Modell." if group == "A" else "hypothetisch (Höhenschema 12,00 m), nicht gesichert."))
    H("4  Gewerbeflächen"); rows = [["Gewerbe"] + names, ["Gastraum (Plätze innen)"] + [f"{d_['stp']['sitze_innen']} / {num(d_['stp']['gastro_m2'])} m²" for _, _, _, d_ in O], ["Gastro-Mietfläche"] + [f"{num(v['gastro_nf'])} m²" for _, v, _, _ in O], ["Büro netto"] + [f"{num(v['buero_nf'],1)} m²" for _, v, _, _ in O]]; TB(rows)
    H("5  Gastro-Konzept"); PR("Große Gastro mit Bar, neu und modern (Olivgrün, Cognac-Samt, Bordeaux, Eiche, schwarzer Stahl). Der Gastro-Umsatz wird in diesem Plan proportional zu den Innenplätzen angesetzt (342 € je Platz und Monat wie im früheren Plan, Marge vor Pacht 30,5 %); das ist eine Annahme ohne Marktbeleg.")
    H("6  Kosten"); labels = [p[0] for p in O[0][2]["posten"]]; rows = [["Posten"] + names] + [[lab] + [eur(r["posten"][i][1]) for _, _, r, _ in O] for i, lab in enumerate(labels)] + [["Gesamtkosten"] + [eur(r["gesamt"]) for _, _, r, _ in O]]; TB(rows, boldrows=(len(rows) - 1,))
    H("7  Finanzierung mit KfW-Krediten"); rows = [["Mittel"] + names + ["Zins"], ["KfW 298 (vermietete Wohnungen)"] + [eur(r["kfw_w"]) for _, _, r, _ in O] + ["2,50 %"], ["KfW 299 (Gastro und Büros)"] + [eur(r["kfw_g"]) for _, _, r, _ in O] + ["3,00 %"], ["Hausbank"] + [eur(r["bank"]) for _, _, r, _ in O] + ["4,00 %"], ["Mischzins"] + [pct(r["mischzins"], 2) for _, _, r, _ in O] + [""], ["Rate pro Monat"] + [eur(r["rate_m"]) for _, _, r, _ in O] + [""]]; TB(rows)
    PR("Voraussetzungen: Effizienzhaus 40, QNG-Zertifikat, Heizung ohne Öl/Gas/Biomasse, Energieeffizienz-Experte, Antrag vor Baubeginn über die Hausbank. KfW-Zinsen und Rahmen sind nicht bestätigt (Quellen 0,6 % bis 3,05 %).")
    H("8  Abzahlungsplan"); rows = [["Jahr"] + [f"{n_}: Zinsen / Tilgung / Restschuld" for n_ in names]] + [[str(y)] + [f"{eur(r['plan'][y-1][1])} / {eur(r['plan'][y-1][2])} / {eur(r['plan'][y-1][3])}" for _, _, r, _ in O] for y in (1, 2, 5, 10, 20, 35)]; TB(rows)
    H("9  Ertrag und Kapitaldienst"); rows = [["Position (pro Jahr)"] + names, ["Kaltmiete Wohnen"] + [eur(r["wohn"]) for _, _, r, _ in O], ["Büromieten"] + [eur(r["buero"]) for _, _, r, _ in O], ["Pacht Gastro"] + [eur(r["pacht"]) for _, _, r, _ in O], ["Rücklagen 10 %"] + [eur(-r["rueck"]) for _, _, r, _ in O], ["Betriebsergebnis Gastro nach Pacht"] + [eur(r["gastro_nach"]) for _, _, r, _ in O], ["Ertrag vor Kapitaldienst"] + [eur(r["vor"]) for _, _, r, _ in O], ["Kapitaldienst"] + [eur(-r["rate_j"]) for _, _, r, _ in O], ["Ergebnis"] + [eur(r["cash"], True) for _, _, r, _ in O]]; TB(rows, boldrows=(6, 8))
    H("10  Ertragswert und Beleihung"); rows = [["Wert"] + names, ["Ertragswert (vereinfacht)"] + [eur(r["ew"]) for _, _, r, _ in O], ["Beleihungswert"] + [eur(r["beleih"]) for _, _, r, _ in O], ["Beleihungsauslauf"] + [pct(r["auslauf"]) for _, _, r, _ in O], ["Betrag über 80 % Auslauf"] + [eur(r["ueber80"]) for _, _, r, _ in O]]; TB(rows)
    PR("Vereinfachte Wertermittlung, kein Gutachten (Liegenschaftszins Wohnen 3,85 %, Büro 5,35 %, Gastro 6,35 %, Restnutzungsdauer 60 Jahre, Bewirtschaftung 13,75 %). Die Bank verlangt zusätzliche Sicherheiten oder Eigenkapital.")
    H("11  Prognose und Belastungstest"); PR("Ergebnis pro Monat nach Kapitaldienst (Plan) und Stresstests: siehe Präsentation, Folie 'Prognose und Belastungstest'. Ohne Gastro fällt das Ergebnis auf " + " / ".join(eur(((r["wohn"] + r["buero"]) * 0.9 - r["rate_j"]) / 12, True) for _, _, r, _ in O) + " pro Monat.")
    H("12  Prüfung: ein 60-m²-Apartment als Airbnb"); PR("Airbnb lohnt sich in Werlte nicht (Auslastung amtlich 17 %, break-even gegenüber langfristiger Miete bei ca. 63 %). Werte aus dem früheren Plan.")
    H("13  Chancen und Risiken"); PR("Chancen: zentrale Lage an der neuen Mitte, KfW-Förderung, Gastro und Wohnen aus einer Hand. Risiken: Beleihungsauslauf, Abhängigkeit von der Gastro, Zinsen nach 10 Jahren, Genehmigung (Baufenster, Dachform, Höhe), Stellplätze" + (", Kragarme und Feuerwehrzufahrt unter dem Vorbau (B)." if group == "B" else "."))
    H("14  Zeitplan und nächste Schritte"); PR("Okt 2026 Gespräch mit der Stadt (Folgetermin nach dem Termin mit Herrn Keve am 06.10.); Nov 2026 Architekt und Baufenster; Dez 2026 Hausbank und KfW; Q1 2027 Bauantrag und KfW-Antrag; Mitte 2027 Baubeginn (15 Monate); Ende 2028 Eröffnung.")
    H("15  Anlagen und Unterlagen für die Bank"); PR("Präsentation Businessplan_Praesentation_" + group + ".pptx, Kosten-Nutzen pro Kopf, Geometrie- und Planunterlagen (studie_gebaeudeformen), Konsistenz- und Planungsrechtsnotizen.")
    out = os.path.join(HERE, f"Businessplan_{group}.docx"); d.save(out); return out
if __name__ == "__main__":
    for g in ("A", "B"): print(pk_pdf(g), plan_docx(g))
