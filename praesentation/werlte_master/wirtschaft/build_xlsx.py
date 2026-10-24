"""Erzeugt Wirtschaftlichkeit_Gesamtkonzept_B3.xlsx (Formeln, Eingaben im Blatt Annahmen). Danach recalc.py und Gegenprobe gegen model.py."""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from params import PARAMS
OUT = sys.argv[1] if len(sys.argv) > 1 else "Wirtschaftlichkeit_Gesamtkonzept_B3.xlsx"
GEO = json.load(open("../geometry/variant_B3.json")); MET = json.load(open("../geometry/variants3_metrics.json"))["K"]
F = "Arial"; base = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True); inp = Font(name=F, size=10, color="0000FF"); note = Font(name=F, size=9, italic=True, color="5B6770")
hdr_font = Font(name=F, size=10, bold=True, color="FFFFFF"); title = Font(name=F, size=16, bold=True, color="2F4A3E"); sub = Font(name=F, size=10, color="5B6770")
hf = PatternFill("solid", fgColor="2F4A3E"); sf = PatternFill("solid", fgColor="E3E8E5"); kf = PatternFill("solid", fgColor="FFF2CC"); rf = PatternFill("solid", fgColor="F8D7D3")
thin = Side(style="thin", color="D0D0D0"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
EUR = '#,##0 "€";-#,##0 "€";"–"'; PCT = "0.0%"; X2 = "0.00"; NUM = "#,##0"; M2 = '#,##0.0 "m²"'
wb = Workbook()
def sheet(name, t, s="", widths=(), first=False):
    ws = wb.active if first else wb.create_sheet(); ws.title = name; ws.sheet_view.showGridLines = False; ws["A1"] = t; ws["A1"].font = title; ws["A2"] = s; ws["A2"].font = sub
    for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
    return ws
def put(ws, ref, v, fmt=None, font=None, fill=None, wrap=False):
    c = ws[ref]; c.value = v; c.font = font or base; c.border = box
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if wrap: c.alignment = Alignment(wrap_text=True, vertical="top")
    return c
def hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t); c.font = hdr_font; c.fill = hf; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = box
# ---------------- Annahmen ----------------
wsA = sheet("Annahmen", "Annahmen", "Alle Eingaben (blau). Spalten D–F: vorsichtig / realistisch / günstig. Stand 22.10.2026. Status je Zeile: Geometrie, Vorgabe, Annahme, Quelle, noch zu bestätigen.", (6, 62, 18, 13, 13, 13, 38, 90), first=False)
wb.remove(wb["Sheet"]) if "Sheet" in wb.sheetnames else None
hdr(wsA, 4, ["", "Parameter", "Einheit", "vorsichtig", "realistisch", "günstig", "Status", "Quelle / Begründung"])
ROW = {}; r = 5
for p in PARAMS:
    if p[0] == "sec":
        for col in range(1, 9): wsA.cell(row=r, column=col).fill = sf
        wsA.cell(row=r, column=2, value=p[1]).font = bold; r += 1; continue
    k, lab, unit, v, rr, g, st, src = p; ROW[k] = r
    put(wsA, f"A{r}", k, font=note); put(wsA, f"B{r}", lab, wrap=True); put(wsA, f"C{r}", unit)
    fmt = "0.00" if isinstance(v, float) and abs(v) < 100 else NUM
    for j, val in enumerate((v, rr, g)): put(wsA, f"{L(4 + j)}{r}", val, fmt, inp)
    put(wsA, f"G{r}", st, font=note, wrap=True); put(wsA, f"H{r}", src, font=note, wrap=True); r += 1
wsA.freeze_panes = "D5"
def PAR(k, sc):  # Referenz auf Annahmen mit Szenario-Index (Zelle)
    return f"INDEX(Annahmen!$D${ROW[k]}:$F${ROW[k]},{sc})"
# ---------------- Rechenkern ----------------
wsK = sheet("Rechenkern", "Rechenkern", "Jede Spalte ist ein vollständiger Rechenfall. Parameter kommen aus Annahmen (Szenario-Nr. 1 vorsichtig, 2 realistisch, 3 günstig); Zeilen 5–11 sind Fall-Einstellungen.", [46, 10] + [14] * 20)
CASES = [("V 20 %", 1, .2), ("R 20 %", 2, .2), ("G 20 %", 3, .2), ("V 0 %", 1, 0), ("V 10 %", 1, .1), ("V 30 %", 1, .3), ("R 0 %", 2, 0), ("R 10 %", 2, .1), ("R 30 %", 2, .3), ("G 0 %", 3, 0), ("G 10 %", 3, .1), ("G 30 %", 3, .3),
         ("R 20 %, Tilgung 1 %", 2, .2, dict(tilg_ov=1)), ("R Baukosten +10 %", 2, .2, dict(bau_f=1.1)), ("R Baukosten +20 %", 2, .2, dict(bau_f=1.2)), ("R Zins +1 pp", 2, .2, dict(dz=1)), ("R Zins +2 pp", 2, .2, dict(dz=2)),
         ("R Leerstand +5 pp", 2, .2, dict(dv=5)), ("R Leerstand +10 pp", 2, .2, dict(dv=10)), ("R Miete −10 %", 2, .2, dict(mf=.9))]
COL = {c[0]: L(3 + i) for i, c in enumerate(CASES)}
hdr(wsK, 4, ["Größe", "Einheit"] + [c[0] for c in CASES])
FORM = [("sc", "Szenario-Nr.", "", None), ("q", "Eigenkapitalquote", "%", None), ("bau_f", "Faktor Baukosten", "", None), ("dz", "Zinszuschlag Hausbank", "pp", None), ("dv", "Leerstandszuschlag", "pp", None), ("mf", "Faktor Mieten", "", None), ("tilg_ov", "Tilgung Override (leer = Annahme)", "%", None),
 ("grund", "Grundstück", "€", "={gr_m2}*{gr_preis}"), ("neben", "Erwerbsnebenkosten (Grunderwerbsteuer, Notar)", "€", "=[grund]*({grest}+{notar})/100"), ("kg100", "KG 100 Grundstück", "€", "=[grund]+[neben]"), ("kg200", "KG 200 Herrichten und Erschließen", "€", "={kg200}"),
 ("bgf", "Bruttogrundfläche gesamt", "m²", "={bgf_eg}+{bgf_l1}+{bgf_l2}"), ("kg300", "KG 300+400 Bauwerk inkl. EH-40-Aufschlag", "€", "=[bgf]*{bau_m2}*[bau_f]*(1+{eh40}/100)"), ("ausbau", "Gewerbe-Ausbau durch Vermieter", "€", "={gew_bgf}*{ausbau}"),
 ("kg500", "KG 500 Außenanlagen, Stellplätze, Terrasse", "€", "={kg500_fl}*{kg500_m2}+{kg500_terr}"), ("basis", "Summe KG 300–500", "€", "=[kg300]+[ausbau]+[kg500]"),
 ("b_arch", "KG 700 Architektur und Objektplanung", "€", "=[basis]*{bnk_arch}/100"), ("b_fach", "KG 700 Fachplaner (Statik, TGA, Brandschutz, Energie, QNG)", "€", "=[basis]*{bnk_fach}/100"), ("b_gen", "KG 700 Genehmigungen, Gutachten, Vermessung", "€", "=[basis]*{bnk_gen}/100"), ("b_son", "KG 700 Bauleitung, Versicherungen, Sonstiges", "€", "=[basis]*{bnk_son}/100"),
 ("bnk", "Summe KG 700 Baunebenkosten", "€", "=[b_arch]+[b_fach]+[b_gen]+[b_son]"), ("reserve", "Reserve", "€", "=([kg200]+[basis])*{reserve}/100"), ("S", "Kosten ohne Finanzierungskosten", "€", "=[kg100]+[kg200]+[basis]+[bnk]+[reserve]"),
 ("zb", "Zinssatz Hausbank", "%", "={z_bank}+[dz]"), ("tl", "anfängliche Tilgung", "%", "=IF([tilg_ov]=\"\",{tilg},[tilg_ov])"), ("a", "Finanzierungskostensatz (Bauzeit, Anlauf, Bearbeitung) je € Kredit", "", "=[zb]/100*({bauzeit}/12*{abruf}/100+{anlauf}/12)+{bearb}/100"),
 ("T", "GESAMTKOSTEN inkl. Finanzierungskosten", "€", "=[S]/(1-(1-[q])*[a])"), ("F", "davon Finanzierungskosten", "€", "=[T]-[S]"), ("D", "Fremdkapital", "€", "=(1-[q])*[T]"), ("EK", "Eigenkapital", "€", "=[q]*[T]"),
 ("zins_j", "Zinsen Jahr 1", "€", "=[D]*[zb]/100"), ("tilg_j", "Tilgung Jahr 1", "€", "=[D]*[tl]/100"), ("ds", "Kapitaldienst jährlich", "€", "=[zins_j]+[tilg_j]"), ("rate_m", "Kreditrate monatlich", "€", "=[ds]/12"), ("zins_m", "Zinsbelastung monatlich", "€", "=[zins_j]/12"),
 ("soll_w", "Sollmiete Wohnen jährlich", "€", "={wfl}*{m_wohn}*[mf]*12"), ("soll_g", "Sollmiete Gastronomie jährlich", "€", "={gastro_nf}*{m_gastro}*[mf]*12"), ("soll_b", "Sollmiete Büros jährlich", "€", "={buero_nf}*{m_buero}*[mf]*12"), ("soll", "Sollmiete gesamt jährlich", "€", "=[soll_w]+[soll_g]+[soll_b]"),
 ("eff_w", "Mieteinnahmen Wohnen nach Leerstand", "€", "=[soll_w]*(1-({v_wohn}+[dv])/100)"), ("eff_g", "Mieteinnahmen Gastronomie nach Ausfall", "€", "=[soll_g]*(1-({v_gastro}+[dv])/100)"), ("eff_b", "Mieteinnahmen Büros nach Leerstand", "€", "=[soll_b]*(1-({v_buero}+[dv])/100)"), ("eff", "Mieteinnahmen gesamt nach Leerstand", "€", "=[eff_w]+[eff_g]+[eff_b]"),
 ("k_w", "nicht umlagefähige Kosten Wohnen", "€", "={wfl}*{ih_wohn}+{we}*{vw_wohn}+{wfl}*{nu_wohn}"), ("k_gfix", "Instandhaltung und sonstige Kosten Gewerbe", "€", "=({gastro_nf}+{buero_nf})*({ih_gew}+{nu_gew})"), ("k_gvw", "Verwaltung Gewerbe", "€", "=([soll_g]+[soll_b])*{vw_gew}/100"), ("kosten", "laufende Kosten gesamt", "€", "=[k_w]+[k_gfix]+[k_gvw]"),
 ("noi", "NETTO-MIETERTRAG (vor Kapitaldienst)", "€", "=[eff]-[kosten]"), ("cf", "Cashflow nach Kapitaldienst jährlich", "€", "=[noi]-[ds]"), ("cf_m", "Cashflow nach Kapitaldienst monatlich", "€", "=[cf]/12"), ("dscr", "KAPITALDIENSTDECKUNGSGRAD", "x", "=[noi]/[ds]"), ("rendite", "Nettoanfangsrendite auf Gesamtkosten", "%", "=[noi]/[T]"),
 ("be_w1", "Break-even Wohnmiete für DSCR 1,0 (übrige Mieten unverändert)", "€/m²", "=(1*[ds]-([noi]-[eff_w]))/({wfl}*12*(1-({v_wohn}+[dv])/100))"), ("be_w12", "Break-even Wohnmiete für DSCR 1,2", "€/m²", "=(1.2*[ds]-([noi]-[eff_w]))/({wfl}*12*(1-({v_wohn}+[dv])/100))"),
 ("be_f1", "Break-even Mietfaktor (alle Mieten) für DSCR 1,0", "x", "=(1*[ds]+[k_w]+[k_gfix])/([eff]-[k_gvw])"), ("be_f12", "Break-even Mietfaktor (alle Mieten) für DSCR 1,2", "x", "=(1.2*[ds]+[k_w]+[k_gfix])/([eff]-[k_gvw])"),
 ("dmax", "Fremdkapital, das bei DSCR 1,2 tragbar wäre", "€", "=[noi]/(1.2*([zb]+[tl])/100)"), ("q_need", "dafür nötige Eigenkapitalquote", "%", "=1-[dmax]/([S]+[dmax]*[a])"), ("ek_need", "dafür nötiges Eigenkapital", "€", "=[S]/(1-(1-[q_need])*[a])-[dmax]"),
 ("w_share", "Wohnanteil der Bruttogeschossfläche (Wohnen + EG-Kern)", "%", "=({bgf_l1}+{bgf_l2}+{bgf_eg}-{gew_bgf})/[bgf]"), ("k297", "KfW 297/298: theoretischer Kredit", "€", "=MIN({we}*{kfw297_we},[w_share]*[D])"), ("k299", "KfW 299: theoretischer Kredit", "€", "=MIN({gew_bgf}*{kfw299_m2},(1-[w_share])*[D])"), ("bank_rest", "Hausbank-Anteil", "€", "=[D]-[k297]-[k299]"),
 ("ds_kfw", "Kapitaldienst mit KfW (theoretisch)", "€", "=[k297]*({z_297}+[tl])/100+[k299]*({z_299}+[tl])/100+[bank_rest]*([zb]+[tl])/100"), ("dscr_kfw", "DSCR mit KfW (theoretisch, nicht bestätigt)", "x", "=[noi]/[ds_kfw]"), ("cf_kfw", "Cashflow mit KfW jährlich", "€", "=[noi]-[ds_kfw]"),
 ("nb_z", "NBank-Zuschuss (Obergrenze 40 %, nicht zugesagt)", "€", "={nb_zusch}/100*MIN([w_share]*[S],{nb_cap}*{wfl})"), ("nb_D", "Fremdkapital nach Zuschuss", "€", "=[D]-[nb_z]"), ("nb_noi", "Netto-Mietertrag bei gebundener Miete", "€", "=[noi]-[eff_w]+{wfl}*{nb_miete}*12*(1-{v_wohn}/100)"), ("nb_ds", "Kapitaldienst nach Zuschuss", "€", "=[nb_D]*([zb]+[tl])/100"), ("nb_dscr", "DSCR mit NBank-Zuschuss (theoretisch)", "x", "=[nb_noi]/[nb_ds]")]
FR = {}; r0 = 5
for i, f in enumerate(FORM): FR[f[0]] = r0 + i
import re
def expand(tpl, col):
    sc = f"{col}${FR['sc']}"
    t = re.sub(r"\{(\w+)\}", lambda m: PAR(m.group(1), sc), tpl)
    t = re.sub(r"\[(\w+)\]", lambda m: f"{col}{FR[m.group(1)]}", t); return t
for i, f in enumerate(FORM):
    key, lab, unit, tpl = f; row = r0 + i; put(wsK, f"A{row}", lab, font=bold if key in ("T", "noi", "dscr", "D") else base); put(wsK, f"B{row}", unit, font=note)
    for ci, c in enumerate(CASES):
        col = L(3 + ci); ov = c[3] if len(c) > 3 else {}
        if key == "sc": put(wsK, f"{col}{row}", c[1], NUM, inp)
        elif key == "q": put(wsK, f"{col}{row}", c[2], "0%", inp)
        elif key in ("bau_f", "mf"): put(wsK, f"{col}{row}", ov.get(key, 1.0), "0.00", inp)
        elif key in ("dz", "dv"): put(wsK, f"{col}{row}", ov.get(key, 0), "0.0", inp)
        elif key == "tilg_ov": put(wsK, f"{col}{row}", ov.get(key), "0.0", inp)
        else:
            fmt = {"€": EUR, "%": "0.00", "x": X2, "m²": M2, "pp": "0.0", "€/m²": "0.00"}.get(unit, "0.0000")
            if key in ("q_need", "rendite", "w_share"): fmt = PCT
            put(wsK, f"{col}{row}", expand(tpl, col), fmt, bold if key in ("T", "noi", "dscr") else base, kf if key in ("T", "noi", "dscr") else None)
wsK.freeze_panes = "C5"
def K(case, key): return f"Rechenkern!{COL[case]}{FR[key]}"
# ---------------- Flächen ----------------
wsF = sheet("Flächen", "Flächen- und Wohnungsübersicht (Geometrie B3, Variante K)", "Quelle: geometry/variant_B3.json. Netto = Brutto × 0,88 (Annahme). Alle Flächen hypothetisch bis zur Freigabe; Balkone und Dachterrassen nicht in der Wohnfläche.", (14, 14, 28, 14, 14, 16, 14, 40))
hdr(wsF, 4, ["Nr.", "Geschoss", "Typ", "Brutto m²", "Netto m²", "Sollmiete €/Monat (realistisch)", "Außenraum", "Bemerkung"])
rooms = [r for r in GEO["variants"]["K"]["rooms"] if r["type"] == "apartment"]
rooms.sort(key=lambda r: (r["level"], round(r["rect"][1]), r["rect"][0])); n = {"L1": 0, "L2": 0}; row = 5
for rr in rooms:
    n[rr["level"]] += 1; pre = "1" if rr["level"] == "L1" else "3"; nr = f"{pre}.{n[rr['level']]:02d}"
    typ = {"WE-L1-V": "1. OG vorn, Studio", "WE-L1-H": "1. OG hinten, Zweizimmer", "WE-L2-V": "3. Geschoss vorn, Zweizimmer", "WE-L2-H": "3. Geschoss hinten, Zweizimmer"}[rr["id"][:7]]
    ext = {"WE-L1-V": "Balkon 2,70 × 1,50 m", "WE-L1-H": "–", "WE-L2-V": "Umlauf 0,80 m", "WE-L2-H": "Dachterrasse 20,7 m²"}[rr["id"][:7]]
    put(wsF, f"A{row}", nr); put(wsF, f"B{row}", "1. OG" if rr["level"] == "L1" else "Staffel"); put(wsF, f"C{row}", typ); put(wsF, f"D{row}", rr["area_gross"], "0.00", inp); put(wsF, f"E{row}", f"=D{row}*0.88", "0.00"); put(wsF, f"F{row}", f"=E{row}*Annahmen!$E${ROW['m_wohn']}", EUR); put(wsF, f"G{row}", ext); put(wsF, f"H{row}", "", font=note); row += 1
put(wsF, f"A{row}", "Summe 16 Wohnungen", font=bold); put(wsF, f"D{row}", f"=SUM(D5:D{row-1})", "0.00", bold); put(wsF, f"E{row}", f"=SUM(E5:E{row-1})", "0.00", bold, kf); put(wsF, f"F{row}", f"=SUM(F5:F{row-1})", EUR, bold); WF_ROW = row
row += 2; put(wsF, f"A{row}", "Nutzungsflächen", font=bold); row += 1
hdr(wsF, row, ["Bereich", "", "", "Brutto m²", "Netto / Mietfläche m²", "", "", "Bemerkung"]); row += 1
uses = [("Gastronomie (Gastraum 280, Küche 60, Lager 27, WC 24, Personal 20, Flur 20, Anlieferung 9)", 440, 440, "ca. 120 Innenplätze möbliert, 44 Terrassenplätze"), ("Büros 1 und 2 (je 32 m² brutto)", 64, 56.3, "ohne Foyer und Büro-WC (15,6 m²)"), ("Büro-Foyer und Büro-WC (Alternative K)", 15.6, 0, "keine Mietfläche"),
        ("Wohnkern EG (Halle, Treppe, Flur, Aufzug, Technik)", 56.4, 0, "Erschließung Wohnen"), ("Wohnungen 1. OG", 660, 468.0, "inklusive Kern, Gang"), ("Wohnungen Staffelgeschoss", 539.06, 377.8, "inklusive Kern, Gang")]
for nm, b, nn, bem in uses:
    put(wsF, f"A{row}", nm); wsF.merge_cells(f"A{row}:C{row}"); put(wsF, f"D{row}", b, "0.0", inp); put(wsF, f"E{row}", nn, "0.0", inp); put(wsF, f"H{row}", bem, font=note); row += 1
put(wsF, f"A{row}", "Bruttogeschossfläche gesamt (EG 576 + 1. OG 660 + Staffel 539,06)", font=bold); wsF.merge_cells(f"A{row}:C{row}"); put(wsF, f"D{row}", f"=Annahmen!E{ROW['bgf_eg']}+Annahmen!E{ROW['bgf_l1']}+Annahmen!E{ROW['bgf_l2']}", "0.0", bold, kf)
row += 2; put(wsF, f"A{row}", "Hinweis: Die Geschossflächen (Auskragung 3,00 m, Staffel 1,50 m) sind hypothetisch; Baufenster, GRZ und Abstände sind ungeklärt (C-30).", font=note)
# ---------------- Einnahmen ----------------
wsE = sheet("Einnahmen", "Einnahmen und laufende Kosten (Eigenkapital 20 %, Szenarien)", "Getrennte Einnahmebereiche. Mieten netto kalt ohne Umsatzsteuer. Gastronomie als Vermietung an einen Betreiber, kein eigener Restaurantbetrieb.", (62, 16, 16, 16, 60))
hdr(wsE, 4, ["Position", "vorsichtig", "realistisch", "günstig", "Hinweis"])
cs = ["V 20 %", "R 20 %", "G 20 %"]
lines = [("A. GASTRONOMIE", None, None), ("Mietfläche m²", "gastro_nf", "P"), ("Miete €/m² Monat", "m_gastro", "P"), ("Sollmiete monatlich", None, "soll_g/12"), ("Sollmiete jährlich", "soll_g", "K"), ("Mietausfall %", "v_gastro", "P"), ("Einnahmen nach Ausfall jährlich", "eff_g", "K"),
         ("B. WOHNUNGEN (16)", None, None), ("Wohnfläche m²", "wfl", "P"), ("Miete €/m² Monat", "m_wohn", "P"), ("Sollmiete monatlich", None, "soll_w/12"), ("Sollmiete jährlich", "soll_w", "K"), ("Leerstand und Mietausfall %", "v_wohn", "P"), ("Einnahmen nach Leerstand jährlich", "eff_w", "K"),
         ("C. BÜROS (2)", None, None), ("vermietbare Nettofläche m²", "buero_nf", "P"), ("Miete €/m² Monat", "m_buero", "P"), ("Sollmiete monatlich", None, "soll_b/12"), ("Sollmiete jährlich", "soll_b", "K"), ("Leerstand %", "v_buero", "P"), ("Einnahmen nach Leerstand jährlich", "eff_b", "K"),
         ("GESAMT", None, None), ("Sollmiete jährlich", "soll", "K"), ("Sollmiete monatlich", None, "soll/12"), ("Einnahmen nach Leerstand jährlich", "eff", "K"), ("nicht umlagefähige Kosten Wohnen (Instandhaltung, Verwaltung, sonstige)", "k_w", "K"), ("Instandhaltung und sonstige Kosten Gewerbe", "k_gfix", "K"), ("Verwaltung Gewerbe", "k_gvw", "K"),
         ("laufende Kosten gesamt", "kosten", "K"), ("NETTO-MIETERTRAG jährlich", "noi", "K"), ("Netto-Mietertrag monatlich", None, "noi/12")]
row = 5
for lab, key, mode in lines:
    if key is None and mode is None:
        for c in range(1, 6): wsE.cell(row=row, column=c).fill = sf
        wsE.cell(row=row, column=1, value=lab).font = bold; row += 1; continue
    put(wsE, f"A{row}", lab, font=bold if "NETTO" in lab else base)
    for j, cs_ in enumerate(cs):
        ref = f"{L(2 + j)}{row}"
        if mode == "P": put(wsE, ref, f"=INDEX(Annahmen!$D${ROW[key]}:$F${ROW[key]},{j + 1})", "0.0", base)
        elif mode == "K": put(wsE, ref, f"={K(cs_, key)}", EUR, bold if "NETTO" in lab else base, kf if "NETTO" in lab else None)
        else: a, b = mode.split("/"); put(wsE, ref, f"={K(cs_, a)}/{b}", EUR)
    if lab.startswith("Miete €/m² Monat") and False: pass
    row += 1
row += 1; put(wsE, f"A{row}", "Kontrolle Gastronomie: Pachtmiete / Umsatz", font=bold); row += 1
put(wsE, f"A{row}", "angenommener Jahresumsatz des Betreibers je Innenplatz (Annahme, nicht belegt) €"); 
for j in range(3): put(wsE, f"{L(2 + j)}{row}", [4000, 5500, 7000][j], EUR, inp)
put(wsE, f"E{row}", "Annahme: ländlicher Restaurantbetrieb, nicht regional belegt", font=note); row += 1
put(wsE, f"A{row}", "Plätze innen + Terrasse (möbliert normal 120 + 44)"); 
for j in range(3): put(wsE, f"{L(2 + j)}{row}", 164, NUM, inp)
row += 1; put(wsE, f"A{row}", "Pacht in % vom Umsatz (Richtwert Branche 8–12 %, Ratgeber)")
for j in range(3): put(wsE, f"{L(2 + j)}{row}", f"={L(2 + j)}{row - 6 - 0 if False else 9}/({L(2 + j)}{row - 2}*{L(2 + j)}{row - 1})", PCT)
E_PACHT_ROW = row
# ---------------- Kosten ----------------
wsC = sheet("Kosten", "Kostenschätzung nach DIN 276 (Eigenkapital 20 %, Szenarien)", "Preise brutto inkl. 19 % MwSt, Vorsteuerabzug nicht angesetzt. Fehlende Kostendaten sind Annahmen (Blatt Annahmen). Ohne Bodengutachten und Lageplan.", (66, 16, 16, 16, 40))
hdr(wsC, 4, ["Kostengruppe", "vorsichtig", "realistisch", "günstig", "Hinweis"])
kl = [("KG 100 Grundstück (960 m² × 230 €/m²)", "grund", ""), ("Erwerbsnebenkosten (Grunderwerbsteuer 5 %, Notar und Grundbuch 2 %)", "neben", ""), ("KG 200 Herrichten und Erschließen", "kg200", "Pauschale, Grundstückszustand unbekannt"), ("KG 300+400 Bauwerk (Bau- und Technik) inkl. Effizienzhaus-40-Aufschlag", "kg300", "BGF 1.775 m² × €/m²"),
      ("Gewerbe-Ausbau durch Vermieter (Anschlüsse, Abluft, Fettabscheider)", "ausbau", "Betreiberküche und Mieterausbau durch Mieter"), ("KG 500 Außenanlagen, Stellplätze, Terrasse mit Vordach", "kg500", "HYPOTHETISCH (32 × 30 m)"), ("KG 700 Architektur und Objektplanung", "b_arch", ""), ("KG 700 Fachplaner: Statik, TGA, Brandschutz, Energieberatung, QNG", "b_fach", ""),
      ("KG 700 Genehmigungen, Gutachten, Vermessung", "b_gen", ""), ("KG 700 Bauleitung, Versicherungen, Sonstiges", "b_son", ""), ("Reserve", "reserve", ""), ("Zwischensumme ohne Finanzierungskosten", "S", ""), ("Finanzierungskosten (Bauzeitzinsen, Anlaufzinsen, Bereitstellung)", "F", "auf Fremdkapital mit Hausbankzins"), ("GESAMTKOSTEN", "T", "")]
row = 5
for lab, key, hint in kl:
    put(wsC, f"A{row}", lab, font=bold if key in ("S", "T") else base)
    for j, cs_ in enumerate(cs): put(wsC, f"{L(2 + j)}{row}", f"={K(cs_, key)}", EUR, bold if key in ("S", "T") else base, kf if key == "T" else None)
    put(wsC, f"E{row}", hint, font=note); row += 1
put(wsC, f"A{row}", "Gesamtkosten je m² BGF"); 
for j, cs_ in enumerate(cs): put(wsC, f"{L(2 + j)}{row}", f"={K(cs_, 'T')}/{K(cs_, 'bgf')}", EUR)
row += 1; put(wsC, f"A{row}", "Gesamtkosten je m² Nutzfläche (Wohnen + Gewerbe netto)")
for j, cs_ in enumerate(cs): put(wsC, f"{L(2 + j)}{row}", f"={K(cs_, 'T')}/(INDEX(Annahmen!$D${ROW['wfl']}:$F${ROW['wfl']},{j + 1})+INDEX(Annahmen!$D${ROW['gastro_nf']}:$F${ROW['gastro_nf']},{j + 1})+INDEX(Annahmen!$D${ROW['buero_nf']}:$F${ROW['buero_nf']},{j + 1}))", EUR)
row += 2; put(wsC, f"A{row}", "Nicht enthalten: Mieterausbau, Küche und Inventar des Betreibers; KG 600; Umsatzsteuer-Option bei Gewerbe (Vorsteuer); Stellplatz-Ablöse; Fördermittel.", font=note)
# ---------------- Finanzierung ----------------
wsN = sheet("Finanzierung", "Finanzierung bei unterschiedlichen Eigenkapitalquoten (ohne Förderung)", "Zinsen und Tilgung im Jahr 1; Hausbank 10 Jahre fest. Eigenkapital = Anteil der Gesamtkosten inkl. Finanzierungskosten.", [44] + [14] * 12)
hdr(wsN, 4, ["Größe"] + ["V 0 %", "V 10 %", "V 20 %", "V 30 %", "R 0 %", "R 10 %", "R 20 %", "R 30 %", "G 0 %", "G 10 %", "G 20 %", "G 30 %"])
nl = [("Gesamtkosten", "T", EUR), ("Eigenkapital", "EK", EUR), ("Fremdkapital", "D", EUR), ("Zinssatz Hausbank %", "zb", "0.00"), ("Zinsbelastung monatlich", "zins_m", EUR), ("anfängliche Tilgung %", "tl", "0.00"), ("Kreditrate monatlich", "rate_m", EUR), ("Kapitaldienst jährlich", "ds", EUR),
      ("Netto-Mietertrag jährlich", "noi", EUR), ("Cashflow nach Kapitaldienst jährlich", "cf", EUR), ("Cashflow nach Kapitaldienst monatlich", "cf_m", EUR), ("Kapitaldienstdeckungsgrad", "dscr", X2), ("Nettoanfangsrendite", "rendite", PCT)]
order = ["V 0 %", "V 10 %", "V 20 %", "V 30 %", "R 0 %", "R 10 %", "R 20 %", "R 30 %", "G 0 %", "G 10 %", "G 20 %", "G 30 %"]
for i, (lab, key, fmt) in enumerate(nl):
    rr = 5 + i; put(wsN, f"A{rr}", lab, font=bold if key == "dscr" else base)
    for j, c in enumerate(order): put(wsN, f"{L(2 + j)}{rr}", f"={K(c, key)}", fmt, bold if key == "dscr" else base, kf if key == "dscr" else None)
rr = 5 + len(nl); put(wsN, f"A{rr}", "Bewertung (rechnerisch)", font=bold)
for j, c in enumerate(order):
    col = L(2 + j); put(wsN, f"{col}{rr}", f'=IF({col}{5 + 11}>=1.2,"tragfähig",IF({col}{5 + 11}>=1,"knapp","nicht tragfähig"))', None, bold)
rr += 2; put(wsN, f"A{rr}", "Kapitaldienst bei 1 % statt 2 % Tilgung (R, EK 20 %)", font=bold); put(wsN, f"B{rr}", f"={K('R 20 %, Tilgung 1 %', 'ds')}", EUR); put(wsN, f"C{rr}", "DSCR"); put(wsN, f"D{rr}", f"={K('R 20 %, Tilgung 1 %', 'dscr')}", X2)
rr += 1; put(wsN, f"A{rr}", "Fremdkapital, das bei DSCR 1,2 tragbar wäre (R, EK 20 %)", font=bold); put(wsN, f"B{rr}", f"={K('R 20 %', 'dmax')}", EUR); put(wsN, f"C{rr}", "nötiges Eigenkapital"); put(wsN, f"D{rr}", f"={K('R 20 %', 'ek_need')}", EUR); put(wsN, f"E{rr}", "Quote"); put(wsN, f"F{rr}", f"={K('R 20 %', 'q_need')}", PCT)
rr += 1; put(wsN, f"A{rr}", "Break-even Wohnmiete DSCR 1,0 / 1,2 (R, EK 20 %) €/m²"); put(wsN, f"B{rr}", f"={K('R 20 %', 'be_w1')}", "0.00"); put(wsN, f"C{rr}", f"={K('R 20 %', 'be_w12')}", "0.00")
rr += 1; put(wsN, f"A{rr}", "Break-even Faktor aller Mieten DSCR 1,0 / 1,2 (R, EK 20 %)"); put(wsN, f"B{rr}", f"={K('R 20 %', 'be_f1')}", X2); put(wsN, f"C{rr}", f"={K('R 20 %', 'be_f12')}", X2)
# ---------------- Sensitivität ----------------
wsS = sheet("Sensitivität", "Sensitivität (realistisches Szenario, Eigenkapital 20 %)", "Jede Spalte ist ein voller Rechenfall.", [44] + [15] * 9)
sl = ["R 20 %", "R Baukosten +10 %", "R Baukosten +20 %", "R Zins +1 pp", "R Zins +2 pp", "R Leerstand +5 pp", "R Leerstand +10 pp", "R Miete −10 %", "R 20 %, Tilgung 1 %"]
hdr(wsS, 4, ["Größe"] + sl)
for i, (lab, key, fmt) in enumerate([("Gesamtkosten", "T", EUR), ("Fremdkapital", "D", EUR), ("Kapitaldienst jährlich", "ds", EUR), ("Netto-Mietertrag jährlich", "noi", EUR), ("Cashflow jährlich", "cf", EUR), ("Kapitaldienstdeckungsgrad", "dscr", X2), ("Break-even Wohnmiete DSCR 1,0 €/m²", "be_w1", "0.00")]):
    rr = 5 + i; put(wsS, f"A{rr}", lab, font=bold if key == "dscr" else base)
    for j, c in enumerate(sl): put(wsS, f"{L(2 + j)}{rr}", f"={K(c, key)}", fmt, bold if key == "dscr" else base, kf if key == "dscr" else None)
put(wsS, "A14", "Bei allen Fällen liegt der Deckungsgrad unter 1,0: Die Aussage hängt nicht an einzelnen Annahmen.", font=note)
# ---------------- Förderung ----------------
wsP = sheet("Förderung", "Förderung: theoretische Kredithöhe und Einfluss auf die Finanzierung (realistisch, Eigenkapital 20 %)", "Unterscheidung: theoretisch möglich / fachlich plausibel / bestätigte Zusage (keine). Konditionen sind zeitabhängig; Zinssätze noch zu bestätigen. Abruf der Quellen 22.10.2026.", (34, 22, 22, 22, 22, 60))
hdr(wsP, 4, ["Programm", "theoretisch möglich", "fachlich plausibel (Einschätzung)", "bestätigte Zusage", "theoretische Kredit-/Zuschusshöhe", "Bedingungen, Hinweise, Kombinationsausschluss"])
prog = [("KfW 297/298 Klimafreundlicher Neubau Wohngebäude", "ja (Wohnanteil ca. 72 % der BGF)", "plausibel, wenn Effizienzhaus 40 bzw. 55 mit QNG erreichbar; Misch-Nutzung zu prüfen", "keine", f"={K('R 20 %', 'k297')}", "bis 100.000 € je Wohneinheit, bis 150.000 € mit QNG-PLUS/-PREMIUM (16 WE: bis 2,4 Mio €); reine Zinsverbilligung ohne Tilgungszuschuss; Laufzeit bis 35 Jahre, Zinsbindung bis 10 Jahre; Effizienzhaus-55-Stufe befristet bis 31.12.2026 und mit begrenzten Bundesmitteln; Antrag vor Baubeginn; Zinssatz noch zu bestätigen"),
        ("KfW 296 Klimafreundlicher Neubau im Niedrigpreissegment", "nur wenn alle Programmkriterien erfüllt", "unwahrscheinlich: Kostengrenzen, Wohnflächenoptimierung und Lebenszykluskriterien, Gewerbeanteil und Dachterrassen nicht geprüft", "keine", "bis 100.000 € je Wohneinheit (16 WE: 1,6 Mio €)", "mindestens Effizienzhaus 55, Treibhausgasgrenze, keine fossile Wärme, Wohnflächenoptimierung nach WoFlV; nicht mit 297/298/299 für dieselbe Einheit kombinierbar; Merkblatt 06/2026 nicht abgerufen"),
        ("KfW 299 Klimafreundlicher Neubau Nichtwohngebäude (Gewerbeteil)", "ja für Gastronomie und Büros (Nichtwohnanteil ca. 28 %)", "unklar: Bezugsfläche, Mischnutzung und Abgrenzung zum Wohnteil sind zu klären", "keine", f"={K('R 20 %', 'k299')}", "laut Drittseiten bis 1.500 €/m² (2.000 mit QNG, 1.000 bei Effizienzgebäude 55); Merkblatt 07/2026 (ab 30.07.2026) nicht abrufbar; 10 Jahre zweckentsprechende Nutzung; Antrag vor Baubeginn; Zinssatz nicht gefunden"),
        ("NBank Mietwohnraumförderung Neubau (Niedersachsen, seit 01.07.2026)", "ja für den Wohnteil (gebundene Mieten)", "nicht plausibel mit den Marktmieten: Mietbindung 35 Jahre, Mietstufe Werlte unbekannt", "keine", f"={K('R 20 %', 'nb_z')}", "Zuschuss bis 40 % der förderfähigen Kosten, bis 80 % Gesamtförderung bei niedrigen Einkommen (Pressemeldung); Obergrenzen je Wohnung nicht abgerufen; Miete 6,30–7,80 €/m² (Stand 04/2026), Effizienzhaus 40, mindestens 3 Wohnungen, Antrag vor Maßnahmenbeginn"),
        ("Kommunale Programme Werlte / Samtgemeinde Werlte", "nicht ermittelt", "nicht beurteilbar", "keine", "–", "bei der Stadt erfragen")]
for i, p in enumerate(prog):
    rr = 5 + i
    for j, v in enumerate(p): put(wsP, f"{L(1 + j)}{rr}", v, EUR if (j == 4 and isinstance(v, str) and v.startswith("=")) else None, base, None, True)
    wsP.row_dimensions[rr].height = 150
rr = 11; put(wsP, f"A{rr}", "Einfluss auf die Finanzierung (theoretisch, nicht zugesagt)", font=bold); rr += 1
hdr(wsP, rr, ["Größe", "ohne Förderung", "mit KfW 297 + 299 (theoretisch)", "mit NBank-Zuschuss (theoretisch)", "", "Bemerkung"]); rr += 1
put(wsP, f"A{rr}", "Fremdkapital"); put(wsP, f"B{rr}", f"={K('R 20 %', 'D')}", EUR); put(wsP, f"C{rr}", f"={K('R 20 %', 'D')}", EUR); put(wsP, f"D{rr}", f"={K('R 20 %', 'nb_D')}", EUR); rr += 1
put(wsP, f"A{rr}", "davon KfW"); put(wsP, f"B{rr}", 0, EUR); put(wsP, f"C{rr}", f"={K('R 20 %', 'k297')}+{K('R 20 %', 'k299')}", EUR); put(wsP, f"D{rr}", 0, EUR); rr += 1
put(wsP, f"A{rr}", "Kapitaldienst jährlich"); put(wsP, f"B{rr}", f"={K('R 20 %', 'ds')}", EUR); put(wsP, f"C{rr}", f"={K('R 20 %', 'ds_kfw')}", EUR); put(wsP, f"D{rr}", f"={K('R 20 %', 'nb_ds')}", EUR); rr += 1
put(wsP, f"A{rr}", "Netto-Mietertrag jährlich"); put(wsP, f"B{rr}", f"={K('R 20 %', 'noi')}", EUR); put(wsP, f"C{rr}", f"={K('R 20 %', 'noi')}", EUR); put(wsP, f"D{rr}", f"={K('R 20 %', 'nb_noi')}", EUR); put(wsP, f"F{rr}", "NBank: gebundene Miete 7,80 €/m² senkt den Ertrag", font=note); rr += 1
put(wsP, f"A{rr}", "Kapitaldienstdeckungsgrad", font=bold); put(wsP, f"B{rr}", f"={K('R 20 %', 'dscr')}", X2, bold, kf); put(wsP, f"C{rr}", f"={K('R 20 %', 'dscr_kfw')}", X2, bold, kf); put(wsP, f"D{rr}", f"={K('R 20 %', 'nb_dscr')}", X2, bold, kf); rr += 1
put(wsP, f"A{rr}", "Zinssatz KfW 297 / 299 (Annahme realistisch, noch zu bestätigen)"); put(wsP, f"B{rr}", f"=Annahmen!E{ROW['z_297']}", "0.00"); put(wsP, f"C{rr}", f"=Annahmen!E{ROW['z_299']}", "0.00"); rr += 2
put(wsP, f"A{rr}", "Bewertung: Auch mit den theoretischen Höchstbeträgen bleibt der Deckungsgrad deutlich unter 1,0. Die Finanzierung wird durch Förderung allein nicht tragfähig; Förderung ist keine zugesagte Größe.", font=bold)
# ---------------- Quellen ----------------
wsQ = sheet("Quellen", "Quellen und Abrufdatum (22.10.2026)", "Aus der Websuche; offizielle Seiten von KfW und NBank waren aus der Arbeitsumgebung nicht direkt abrufbar (Netzwerksperre), Merkblätter wurden nicht im Volltext gelesen.", (34, 70, 90))
hdr(wsQ, 4, ["Thema", "Quelle", "Aussage und Vorbehalt"])
qs = [("KfW 297/298", "https://baukostenplan.de/ratgeber/kfw-297-298-klimafreundlicher-neubau-2026; https://baumentor.de/klimafreundlicher-neubau-kfw-programm-neubaufoerderung/; https://hausbaujournal.com/hausbau-planen/kfw-297-298-klimafreundlicher-neubau/", "100.000 € je WE, 150.000 € mit QNG-PLUS/-PREMIUM; Zinsangaben widersprüchlich (0,6 % März 2026; 1,96 % laut Quelle vom 01.10.2026; 3,05 % EH 40 laut Drittseite, KfW 29.09.2026); Effizienzhaus-55-Stufe bis 31.12.2026. Zinsen noch zu bestätigen."),
      ("KfW 296", "https://www.kfw.de/PDF/Download-Center/F%C3%B6rderprogramme-(Inlandsf%C3%B6rderung)/PDF-Dokumente/6000005195_M_296.pdf; https://baukostenplan.de/ratgeber/kfw-296-klimafreundlicher-neubau-niedrigpreissegment", "bis 100.000 € je WE, mindestens EH 55, Wohnflächenoptimierung, nicht kombinierbar mit 297/298/299 je Einheit. Merkblatt nicht im Volltext gelesen."),
      ("KfW 299", "https://www.kfw.de/PDF/Download-Center/F%C3%B6rderprogramme-(Inlandsf%C3%B6rderung)/PDF-Dokumente/6000005052_M_299.pdf; https://www.fensterversand.com/foerderung/kfw/299.php", "Höchstbeträge je m² laut Drittseite 1.000/1.500/2.000 €; Merkblatt Fassung 07/2026; Zinssatz nicht gefunden."),
      ("NBank Wohnraumförderung", "https://www.nbank.de/Service/Presse/Niedersachsen-stellt-soziale-Wohnraumf%C3%B6rderung-neu-auf/; https://www.nbank.de/medien/nb-media/Downloads/Programminformation/Produktinformationen/Produktinformation-Energetisch-hochwertiger-Mietwohnungsneubau-in-Mehrfamilienh%C3%A4usern.pdf", "Neufassung seit 01.07.2026, Zuschuss bis 40 % der förderfähigen Kosten; Bindung 35 Jahre; Mieten und Kostengrenzen Stand 04/2026 evtl. überholt."),
      ("Wohnungsmieten Werlte/Emsland", "https://www.kleinanzeigen.de/s-werlte-/2-zimmer-wohnung/k0l13476; https://www.immowelt.de/suche/mieten/wohnung/niedersachsen/werlte-49757/ad08de1969; https://www.immobilienscout24.de/immobilienpreise/niedersachsen/emsland-kreis/mietspiegel; https://www.miet-check.de/mietspiegel/kreis/emsland/; https://www.miete-aktuell.de/mietspiegel/Werlte/Werlte/", "Angebote Neubau 9,5–12,5 €/m²; Durchschnitt Werlte 7,78–8,43 €/m²; Stichprobe klein, Angebotsmieten."),
      ("Büromieten", "https://www.immowelt.de/suche/mieten/praxis-bueros/niedersachsen/werlte-49757/ad08de1969; https://www.immobilienscout24.de/gewerbe-flaechen/de/niedersachsen/emsland-kreis/buero-mieten/", "Werlte 9,6–12 €/m² (kleine Einheiten), Meppen 9–9,5, Sögel 6,3–7,8; Angebotspreise."),
      ("Gastronomie-Pacht", "https://gastro-master.de/de/blog/pachtkalkulation-restaurant-berechnen; https://kanzlei-herfurtner.de/gewerbemiete-umsatzmiete-mindestmiete/", "Mietquote 8–12 % vom Nettoumsatz als Richtwert (Ratgeber); keine regionalen Gastro-Flächenmieten gefunden."),
      ("Baukosten", "https://baukostenplan.de/haustypen/mehrfamilienhaus; https://baukostenplan.de/baukosten/niedersachsen", "MFH ca. 2.500–3.500 €/m² Wohnfläche (Ratgeber, nicht BKI); BKI-Regionalfaktor Niedersachsen nicht geprüft."),
      ("Zinsen", "https://www.baufi24.de/presse/pressemitteilungen/zinskommentar-102026/; https://www.sparkasse.de/pk/ratgeber/wohnen/immobilie-erwerben/bauzinsen.html; https://www.interhyp.de/zinsen/", "10 Jahre fest Oktober 2026 ca. 4,0–4,5 %; Durchschnitt Baufi24 September 4,16 %.")]
for i, q in enumerate(qs):
    for j, v in enumerate(q): put(wsQ, f"{L(1 + j)}{5 + i}", v, None, base, None, True)
    wsQ.row_dimensions[5 + i].height = 80
# ---------------- Prüfung ----------------
wsV = sheet("Prüfung", "Plausibilitäts- und Konsistenzprüfung", "Alle Prüfungen müssen WAHR sein.", (70, 20, 20, 12))
hdr(wsV, 4, ["Prüfung", "Wert A", "Wert B", "OK"])
checks = [("Gesamtkosten = Σ Kostengruppen (R 20 %)", f"={K('R 20 %', 'T')}", f"={K('R 20 %', 'kg100')}+{K('R 20 %', 'kg200')}+{K('R 20 %', 'basis')}+{K('R 20 %', 'bnk')}+{K('R 20 %', 'reserve')}+{K('R 20 %', 'F')}"),
          ("Eigenkapital + Fremdkapital = Gesamtkosten (R 20 %)", f"={K('R 20 %', 'EK')}+{K('R 20 %', 'D')}", f"={K('R 20 %', 'T')}"), ("Wohnfläche Flächenblatt = Annahme", f"=Flächen!E{WF_ROW}", f"=Annahmen!E{ROW['wfl']}"),
          ("Kapitaldienst = Zinsen + Tilgung (R 20 %)", f"={K('R 20 %', 'ds')}", f"={K('R 20 %', 'zins_j')}+{K('R 20 %', 'tilg_j')}"), ("Netto-Mietertrag = Einnahmen − Kosten (R 20 %)", f"={K('R 20 %', 'noi')}", f"={K('R 20 %', 'eff')}-{K('R 20 %', 'kosten')}"),
          ("Break-even Wohnmiete liefert DSCR 1,0 (R 20 %)", f"={K('R 20 %', 'noi')}-{K('R 20 %', 'eff_w')}+Annahmen!E{ROW['wfl']}*{K('R 20 %', 'be_w1')}*12*(1-Annahmen!E{ROW['v_wohn']}/100)", f"={K('R 20 %', 'ds')}")]
for i, (lab, a, b) in enumerate(checks):
    rr = 5 + i; put(wsV, f"A{rr}", lab); put(wsV, f"B{rr}", a, "#,##0.00"); put(wsV, f"C{rr}", b, "#,##0.00"); put(wsV, f"D{rr}", f"=ABS(B{rr}-C{rr})<0.01", None, bold)
put(wsV, "A12", "Alle Prüfungen bestanden", font=bold); put(wsV, "D12", f"=AND(D5:D{4 + len(checks)})", None, bold, kf)
# ---------------- Übersicht ----------------
wsU = sheet("Übersicht", "Wirtschaftlichkeit Gesamtkonzept Neue Mitte Werlte, Baufeld 3 (Arbeitsgrundlage B3)", "Ergebnis der Prüfung: ohne Förderung in keinem Szenario tragfähig. Alle Werte Annahmen; keine Förderzusage; Architektur, Wirtschaftlichkeit und baurechtliche Machbarkeit getrennt zu bewerten.", (52, 16, 16, 16, 60), first=False)
hdr(wsU, 4, ["Größe (Eigenkapital 20 %)", "vorsichtig", "realistisch", "günstig", "Hinweis"])
ul = [("Gesamtkosten", "T", EUR), ("Eigenkapital", "EK", EUR), ("Fremdkapital", "D", EUR), ("Zinsbelastung monatlich", "zins_m", EUR), ("Kreditrate monatlich (Zins + 2 % Tilgung)", "rate_m", EUR), ("Kapitaldienst jährlich", "ds", EUR), ("Mieteinnahmen nach Leerstand jährlich", "eff", EUR), ("laufende Kosten jährlich", "kosten", EUR),
      ("Netto-Mietertrag jährlich", "noi", EUR), ("Cashflow nach Kapitaldienst jährlich", "cf", EUR), ("Kapitaldienstdeckungsgrad", "dscr", X2), ("Nettoanfangsrendite", "rendite", PCT), ("Break-even Wohnmiete für DSCR 1,0 (€/m²)", "be_w1", "0.00"), ("Break-even Faktor aller Mieten für DSCR 1,0", "be_f1", X2),
      ("nötiges Eigenkapital für DSCR 1,2", "ek_need", EUR), ("DSCR mit KfW 297 + 299 (theoretisch, nicht zugesagt)", "dscr_kfw", X2)]
for i, (lab, key, fmt) in enumerate(ul):
    rr = 5 + i; put(wsU, f"A{rr}", lab, font=bold if key in ("dscr", "noi", "T") else base)
    for j, c in enumerate(cs): put(wsU, f"{L(2 + j)}{rr}", f"={K(c, key)}", fmt, bold if key in ("dscr", "noi", "T") else base, kf if key == "dscr" else None)
rr = 5 + len(ul); put(wsU, f"A{rr}", "Bewertung (rechnerisch)", font=bold)
for j in range(3): col = L(2 + j); put(wsU, f"{col}{rr}", f'=IF({col}15>=1.2,"tragfähig",IF({col}15>=1,"knapp","NICHT tragfähig"))', None, bold, rf)
put(wsU, f"E{rr}", "DSCR ab 1,2 gilt üblicherweise als tragfähig (Annahme); hier weit darunter.", font=note)
wsU[f"A{rr + 2}"] = "Konsistenzprüfung (Blatt Prüfung):"; wsU[f"A{rr + 2}"].font = bold; put(wsU, f"B{rr + 2}", "=Prüfung!D12", None, bold)
order_sheets = ["Übersicht", "Annahmen", "Flächen", "Einnahmen", "Kosten", "Finanzierung", "Sensitivität", "Förderung", "Quellen", "Rechenkern", "Prüfung"]
wb._sheets = [wb[n] for n in order_sheets]
for ws in wb.worksheets: ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr = __import__("openpyxl").worksheet.properties.PageSetupProperties(fitToPage=True); ws.page_setup.fitToHeight = 0
wb.save(OUT); print("gespeichert", OUT, "Zeilen Rechenkern", len(FORM))
json.dump(dict(FR=FR, COL=COL, ROW=ROW), open("xlsx_map.json", "w"))
