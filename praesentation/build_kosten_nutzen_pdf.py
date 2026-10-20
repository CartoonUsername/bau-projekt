import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

B = json.load(open("bank_values.json"))
GREEN = RGBColor(0x2E, 0x4A, 0x3E); RED = RGBColor(0x8B, 0x3A, 0x2F); GREY = RGBColor(0x5A, 0x5F, 0x66)
def n(x, d=0): return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
def eur(x): return ("−" if x < 0 else "") + n(abs(x)) + " €"
def eurs(x): return ("+" if x > 0 else "−" if x < 0 else "") + n(abs(x)) + " €"

kosten = B["k_gesamt"]
G = B["gs_kopf"]; Z = B["szenarien"]
roh = B["rohertrag"]; rueck = -B["c_rueck"] * 12
netto = B["c_netto_immo"] * 12
rate = -B["c_kd"] * 12
cash = B["c_ergebnis"] * 12
betreiber = B["c_gastro"] * 12
zins = B["zins_jahr1"]; tilg = B["tilg_jahr1"]

doc = Document()
sec = doc.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7)
for a in ("left_margin", "right_margin"): setattr(sec, a, Cm(2))
sec.top_margin = Cm(1.4); sec.bottom_margin = Cm(1.2)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

def para(t, bold=False, size=None, color=None, after=4, italic=False, font=None):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if font: r.font.name = font
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
    return p
def head(t): 
    p = para(t, size=14, color=RED, after=3, font="Georgia"); p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True
def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexc); tcPr.append(s)
def table(rows, widths, right=(), boldrows=()):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.autofit = False
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            p = c.paragraphs[0]; r = p.add_run(str(v)); r.font.size = Pt(9.5)
            p.paragraph_format.space_after = Pt(1)
            if j in right: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if i == 0: r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); shade(c, "2E4A3E")
            elif i in boldrows: r.bold = True; shade(c, "E4EDE6")
    for j, w in enumerate(widths): t.columns[j].width = Cm(w)
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

para("NEUE MITTE WERLTE · BAUFELD 3", bold=True, size=9, color=RGBColor(0xB8, 0x8A, 0x2E), after=0)
para("Kosten, Nutzen und Gewinn pro Kopf", size=22, color=GREEN, after=2, font="Georgia")
para("Kurze Erklärung für die Gesellschafter einer Vermietungs-GbR. Basis: Excel-Rechenmodell (12 Wohnungen, 2 Büros, Restaurant und Bar 510 m² als getrennte Betreiberfirma), Jahr 1 nach der Anlaufzeit, ohne Eigenkapital. Alle Zahlen sind Planungsannahmen, Stand Oktober 2026.", italic=True, color=GREY, size=9.5, after=6)

head("1 Was kostet das Projekt?")
para(f"Die Gesamtkosten liegen bei {eur(kosten)} (Grundstück, Bau, Nebenkosten, Ausbau, Mehrkosten für 4 zusätzliche Wohnungen, Finanzierungskosten der Bauzeit). Finanziert wird mit KfW-Krediten ({eur(B['kredit_297'] + B['kredit_298'] + B['kredit_299'])}, {n(B['kredit_kfw_anteil'] * 100, 0)} %) und Hausbank, die Gesellschafter zahlen im Basisfall kein Eigenkapital. Die KfW-Förderung spart {eur(B['foerder_zins_jahr'])} Zinsen im Jahr gegenüber reiner Hausbank-Finanzierung. Zins und 1 % Tilgung betragen {eur(rate / 12)} im Monat ({eur(rate)} im Jahr).")

head("2 Was kommt rein? (Vermietungs-GbR, pro Jahr)")
table([
    ["Position", "pro Jahr"],
    ["Mieten Wohnen (12 Wohnungen, 655 m²) + Büros (2 Ankermieter) + Gastro-Pacht (10 % vom Umsatz)", eur(roh)],
    ["abzüglich Rücklagen (10 %: Verwaltung, Instandhaltung, Leerstand)", "−" + eur(rueck)],
    ["Nettoertrag der GbR", eur(netto)],
    ["Zins und Tilgung an die Banken", "−" + eur(rate)],
    ["Cash der GbR für die Gesellschafter", eurs(cash)],
], [13.2, 3.8], right=(1,), boldrows=(3, 5))
para(f"Nutzen-Faktor: Der Nettoertrag entspricht {n(netto / kosten * 100, 1)} % der Gesamtkosten (Anfangsrendite), die Kosten sind das {n(kosten / netto, 1)}-Fache des Jahresertrags. Die Kapitaldienstdeckung der GbR beträgt {n(B['dscr'], 2)}, Banken verlangen meist 1,2. Das Gastro-Betriebsergebnis ({eur(betreiber)} im Jahr nach Pacht) gehört der Betreiberfirma des Bauherrn und ist nicht enthalten.", after=4)

head("3 Was bekommt jeder? (gleiche Anteile)")
rows = [["Personen", "Cash pro Person / Monat", "Cash pro Person / Jahr", "Steuerlicher Verlust / Jahr", "Nach Steuer / Monat", "Mit Förder-option / Monat"]]
for c, k in ((1, 1), (2, 2), (3, 4), (4, 8)):
    rows.append([str(k), eurs(G[2][c]), eurs(G[1][c]), eurs(G[3][c]), eurs(G[6][c]), eurs(G[9][c])])
table(rows, [1.8, 3.0, 3.0, 3.2, 3.0, 3.0], right=(1, 2, 3, 4, 5))
para(f"Ohne Eigenkapital liegt der Cash knapp unter null. Nach Steuer (35 % Grenzsteuersatz, Annahme) bleibt es positiv, weil Zinsen ({eur(zins)}) und Gebäudeabschreibung 3 % ({eur(B['gs_afa'])}) einen steuerlichen Verlust von {eur(-B['gs_gewinn'])} erzeugen, den die Gesellschafter mit anderen Einkünften verrechnen. Die Tilgung ({eur(tilg)} im Jahr, {eur(tilg / 8)} pro Person bei 8) baut Vermögen auf. Die Förderoption (degressive AfA und Sonder-AfA im Wohnteil) erhöht die Abschreibung auf {eur(B['gs_afa_opt'])}, Fristen müssen geprüft werden.", after=4)
z = Z[2]
para(f"Mit Eigenkapital: Bringen die Gesellschafter {eur(z['ek'])} ein ({eur(z['kopf']['8']['ek'])} pro Kopf bei 8 Personen), erreicht die GbR Deckung 1,2 und {eurs(z['ergebnis_m'])} Cash im Monat ({eurs(z['kopf']['8']['cash_m'])} pro Kopf bei 8 Personen vor Steuern). Mit 8 statt 12 Wohnungen wären es {eur(Z[3]['ek'])} Eigenkapital.", after=4)

head("4 Wo liegt das Risiko?")
for t in [
    f"Die Pacht hängt am Gastro-Umsatz (Plan {eur(B['g_umsatz_monat'])} im Monat). Fällt die Gastro aus, fehlen der GbR zusätzlich {eur(-B['sens'][8][3] * 12)} im Jahr (Deckung {n(B['sens'][8][2], 2)}).",
    f"Der Beleihungsauslauf liegt bei {n(B['auslauf'] * 100, 0)} %: Die Bank verlangt zusätzliche Sicherheiten oder Eigenkapital, die Lücke zur 80-%-Grenze beträgt etwa {n(B['eigenkapital_80'] / 1e6, 2)} Mio € insgesamt.",
    "Eine GbR haftet mit dem persönlichen Vermögen aller Gesellschafter. Die Gastro läuft deshalb in einer getrennten Betreiberfirma, die Pacht zahlt.",
    "Lohn des Bauherrn, Mehrkosten kleinerer Wohnungen über die angesetzten 15.000 € hinaus und Gewerbesteuer sind nicht gerechnet. Steuerberater und Bank müssen das prüfen.",
]:
    p = doc.add_paragraph(style="List Bullet"); r = p.add_run(t); r.font.size = Pt(9.5); p.paragraph_format.space_after = Pt(1)

para("Quelle: Bankunterlagen_Baufeld3.xlsx (Blätter Ertrag und Cashflow, Gesellschafter, Szenarien).", italic=True, color=GREY, size=8.5, after=0)
doc.save("Kosten-Nutzen_pro-Kopf_Baufeld3.docx")
print(cash)
